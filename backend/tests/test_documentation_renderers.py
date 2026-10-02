"""Contratos y documentos OOXML generados desde datos, sin plantillas personales."""

import io
import zipfile
from copy import deepcopy
from datetime import date
from xml.etree import ElementTree as ET

import pytest

from app.services.documentation_catalog import COMMON_FIELDS, DOCUMENT_DEFINITIONS
from app.services.documentation_renderers import render_document

W = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
A = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
EXPECTED_KINDS = {"formulacion_proyecto", "presentacion_proyecto", "acta_inicio", "producto_resultado",
                  "poster_producto", "informe_bimensual", "acta_cierre", "informe_final", "registro_evidencias"}


def sample_field(field):
    if field["type"] == "rows":
        return [{column["key"]: sample_field(column) for column in field["columns"]}]
    if field["type"] == "date":
        return "2026-10-01"
    if field["type"] == "number":
        return 10996585.25 if field.get("unit") == "COP" else 7
    if field["type"] == "select":
        return field["options"][0]["value"]
    return f"Contenido verificable de {field['label']}"


@pytest.fixture
def payload():
    context = {"nombre": "Proyecto sintético para validar documentos", "codigo_sgps": "SGPS-EJEMPLO",
               "objetivo_general": "Medir resultados con evidencias verificables", "objetivos_especificos": ["Identificar necesidades", "Comparar resultados"],
               "descripcion": "Proyecto de prueba con datos sintéticos", "vigencia": 12, "presupuesto_total": 10996585.25,
               "grupo": "Grupo sintético", "semillero": "Semillero sintético", "version": 3, "periodo_bimestre": 2,
               "producto": {"nombre": "Producto sintético", "tipo": "A1", "descripcion": "Resultado medible"}}
    common = {field["key"]: sample_field(field) for field in COMMON_FIELDS}
    return context, common


def data_for(kind):
    return {field["key"]: sample_field(field) for field in DOCUMENT_DEFINITIONS[kind]["fields"]}


def texts(content, fmt):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        parts = ["word/document.xml"] if fmt == "docx" else sorted(
            name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
        namespace = W if fmt == "docx" else A
        tag = ".//w:t" if fmt == "docx" else ".//a:t"
        return "\n".join(node.text or "" for part in parts for node in ET.fromstring(archive.read(part)).findall(tag, namespace))


def test_catalog_has_typed_guided_fields_and_all_nine_document_kinds():
    assert set(DOCUMENT_DEFINITIONS) == EXPECTED_KINDS
    for field in list(COMMON_FIELDS) + [field for definition in DOCUMENT_DEFINITIONS.values() for field in definition["fields"]]:
        assert field["key"] and field["label"] and field["help"]
        assert field["type"] in {"text", "textarea", "date", "number", "rows", "select"}
        assert isinstance(field["required"], bool)
        if field["type"] == "rows":
            assert field["columns"]
            assert all(column["type"] != "rows" for column in field["columns"])
        assert "default" not in field
    assert DOCUMENT_DEFINITIONS["acta_cierre"]["fields"][0]["key"] == "tipo_cierre"
    options = DOCUMENT_DEFINITIONS["acta_cierre"]["fields"][0]["options"]
    assert {option["value"] for option in options} == {"final", "parcial"}
    assert all(key in {field["key"] for field in COMMON_FIELDS} for definition in DOCUMENT_DEFINITIONS.values() for key in definition["required_common"])


@pytest.mark.parametrize("kind", sorted(EXPECTED_KINDS))
def test_each_kind_generates_real_editable_package_with_context_and_draft_status(payload, kind):
    from docx import Document

    context, common = payload
    before = deepcopy((context, common))
    content, mime = render_document(kind, context, common, data_for(kind))
    definition = DOCUMENT_DEFINITIONS[kind]
    assert zipfile.is_zipfile(io.BytesIO(content))
    assert mime == ("application/vnd.openxmlformats-officedocument.wordprocessingml.document" if definition["format"] == "docx" else "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    text = texts(content, definition["format"])
    assert context["nombre"] in text
    assert "SGPS-EJEMPLO" in text
    if kind == "informe_final":
        assert "Borrador para revisión" in Document(io.BytesIO(content)).sections[0].header.paragraphs[0].text
        assert "Versión 3" in Document(io.BytesIO(content)).sections[0].header.paragraphs[0].text
    else:
        assert "Borrador para revisión" in text
        assert "Versión 3" in text
    assert "Paige" not in text and "Juan Ernesto" not in text
    assert before == (context, common)


def test_start_act_preserves_people_activities_budget_and_unsigned_attendees(payload):
    context, common = payload
    common["cronograma"][0]["fecha_textual"] = "Mes 1 - Mes 2"
    data = data_for("acta_inicio")
    content, _ = render_document("acta_inicio", context, common, data)
    text = texts(content, "docx")
    assert "Personal vinculado" in text
    assert "Descripción de actividades" in text
    assert "Presupuesto y cronograma de ejecución" in text
    assert "Mes 1 - Mes 2" in text
    assert "01/10/2026" in text
    assert "COP $ 10.996.585,25" in text
    from docx import Document
    document = Document(io.BytesIO(content))
    attendee_table = next(table for table in document.tables if any(cell.text == "Firma" for cell in table.rows[0].cells))
    assert all(row.cells[-1].text == "" for row in attendee_table.rows[1:])
    assert len(document.tables) >= 5


def test_partial_closure_does_not_assert_final_approval(payload):
    context, common = payload
    data = data_for("acta_cierre")
    data["tipo_cierre"] = "parcial"
    content, _ = render_document("acta_cierre", context, common, data)
    text = texts(content, "docx")
    assert "Cierre parcial del período" in text
    assert "Balance presupuestal" in text
    assert "Evaluación de actividades y entregables" in text
    assert "Proyecto aprobado" not in text
    assert "Firmado" not in text


def test_report_exposes_quantified_results_evidence_and_bimonthly_period(payload):
    context, common = payload
    content, _ = render_document("informe_bimensual", context, common, data_for("informe_bimensual"))
    text = texts(content, "docx")
    assert "Período bimestral 2" in text
    for label in ["Resultados alcanzados", "Indicador", "Meta", "Logro", "Unidad", "Evidencia"]:
        assert label in text
    assert "Referencias" in text


def test_docx_preserves_long_paragraphs_and_grows_tables(payload):
    context, common = payload
    data = data_for("informe_final")
    data["discusion"] = "\n\n".join(f"Párrafo {index} con hallazgos completos y evidencia verificable." for index in range(40))
    data["resultados"] = [dict(data["resultados"][0], entidad=f"Entidad {index}") for index in range(70)]
    content, _ = render_document("informe_final", context, common, data)
    from docx import Document
    document = Document(io.BytesIO(content))
    assert sum(paragraph.text.count("Párrafo") for paragraph in document.paragraphs) == 40
    assert any(len(table.rows) == 71 for table in document.tables)
    assert "Párrafo 39" in texts(content, "docx")
    xml = zipfile.ZipFile(io.BytesIO(content)).read("word/document.xml")
    assert b"tblHeader" in xml


def test_final_report_uses_gcdtp_v01_template_and_fills_every_report_section(payload):
    from docx import Document

    context, common = payload
    context = context | {
        "nombre": "Proyecto de validación GCDTP",
        "objetivo_general": "Objetivo general de validación",
        "objetivos_especificos": ["Objetivo específico primero", "Objetivo específico segundo"],
    }
    common = common | {
        "codigo_cap": "CAP-SINTÉTICO",
        "centro": "Centro de validación",
        "regional": "Regional de validación",
        "ciudad": "Municipio de validación",
        "responsable": "Responsable sintético",
    }
    data = data_for("informe_final") | {
        "autor_informe": "Autora sintética",
        "fecha_entrega": "2026-10-01",
        "clasificacion_informacion": "publica_clasificada",
        "codigo_idea": "IDEA-SINTÉTICA",
        "experto_proyecto": "Experto sintético",
        "linea_tecnologica": "Línea tecnológica sintética",
        "trl_inicial": 2,
        "trl_alcanzado": 5,
        "tecnoparque": "TecnoParque sintético",
        "introduccion": "MARCADOR INTRODUCCIÓN",
        "planteamiento_problema": "MARCADOR PROBLEMA",
        "estado_arte_tecnica": "MARCADOR ESTADO DEL ARTE",
        "metodologia": "MARCADOR METODOLOGÍA",
        "desarrollo_proyecto": "MARCADOR DESARROLLO",
        "resultados": [{
            "entidad": "Población sintética", "actividades_ejecutadas": "Actividad sintética",
            "resultados_alcanzados": "Logro sintético", "indicador": "Indicador sintético",
            "meta": 4, "logro": 3, "unidad": "prototipos", "evidencia": "Soporte sintético",
        }],
        "viabilidad_tecnica": "MARCADOR VIABILIDAD TÉCNICA",
        "viabilidad_operativa": "MARCADOR VIABILIDAD OPERATIVA",
        "viabilidad_economica": "MARCADOR VIABILIDAD ECONÓMICA",
        "viabilidad_normativa": "MARCADOR VIABILIDAD NORMATIVA",
        "viabilidad_mercado": "MARCADOR VIABILIDAD MERCADO",
        "propiedad_intelectual_transferencia": "MARCADOR PROPIEDAD INTELECTUAL",
        "impacto_proyecto": "MARCADOR IMPACTO",
        "conclusiones": "MARCADOR CONCLUSIONES",
        "referencias": "MARCADOR REFERENCIAS",
        "anexos": "MARCADOR ANEXOS",
        "cumplimiento_objetivos": "MARCADOR CUMPLIMIENTO",
        "balance_final": "MARCADOR BALANCE FINAL",
    }

    content, mime = render_document("informe_final", context, common, data)
    document = Document(io.BytesIO(content))
    text = "\n".join([paragraph.text for paragraph in document.paragraphs]
                     + [cell.text for table in document.tables for row in table.rows for cell in row.cells])

    assert mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert document.core_properties.title == "GCDTP-F-023 V01 - Informe final de proyecto"
    assert len(document.tables) == 3
    assert "Sistema Integrado de Gestión y Autocontrol" in text
    assert "INFORME FINAL" in text
    assert "Pública Clasificada (X)" in text
    assert "Autora sintética" in text and "01/10/2026" in text
    assert "IDEA-SINTÉTICA" in text and "Experto sintético" in text
    assert "TecnoParque sintético" in text and "TRL inicial" in text
    assert "Centro de validación" in text and "CAP-SINTÉTICO" in text
    assert "Borrador para revisión" in document.sections[0].header.paragraphs[0].text
    assert "Versión 3" in document.sections[0].header.paragraphs[0].text
    assert "Objetivo general de validación" in text
    assert "Objetivo específico primero" in text and "Objetivo específico segundo" in text
    for marker in [
        "MARCADOR INTRODUCCIÓN", "MARCADOR PROBLEMA", "MARCADOR ESTADO DEL ARTE", "MARCADOR METODOLOGÍA",
        "MARCADOR DESARROLLO", "Población sintética", "Actividad sintética", "Logro sintético",
        "Indicador sintético", "Soporte sintético", "MARCADOR VIABILIDAD TÉCNICA", "MARCADOR VIABILIDAD OPERATIVA",
        "MARCADOR VIABILIDAD ECONÓMICA", "MARCADOR VIABILIDAD NORMATIVA", "MARCADOR VIABILIDAD MERCADO",
        "MARCADOR PROPIEDAD INTELECTUAL", "MARCADOR IMPACTO", "MARCADOR CONCLUSIONES", "MARCADOR REFERENCIAS",
        "MARCADOR ANEXOS", "MARCADOR CUMPLIMIENTO", "MARCADOR BALANCE FINAL",
    ]:
        assert marker in text
    assert "INSTRUCCIONES" not in text
    assert "Control de cambios" not in text
    assert "Describa brevemente" not in text
    assert "Describa los propósitos" not in text
    state_heading = next(paragraph for paragraph in document.paragraphs
                         if paragraph.text.startswith("5. Estado del arte"))
    assert state_heading.paragraph_format.keep_with_next


def test_final_report_includes_real_photos_as_editable_annexes(payload):
    from PIL import Image
    from docx import Document

    context, common = payload
    photo = io.BytesIO()
    Image.new("RGB", (120, 80), "navy").save(photo, format="PNG")
    content, _ = render_document(
        "informe_final", context, common, data_for("informe_final"),
        images=[{"content": photo.getvalue(), "caption": "Evidencia fotográfica sintética"}],
    )
    document = Document(io.BytesIO(content))
    assert len(document.inline_shapes) == 1
    assert any(paragraph.text == "Registro fotográfico" for paragraph in document.paragraphs)
    assert any(paragraph.text == "Evidencia fotográfica sintética" for paragraph in document.paragraphs)


def test_presentation_and_poster_expand_without_truncating_text_or_rows(payload):
    from pptx import Presentation
    context, common = payload
    for kind in ["presentacion_proyecto", "poster_producto"]:
        data = data_for(kind)
        key = next(field["key"] for field in DOCUMENT_DEFINITIONS[kind]["fields"] if field["type"] == "textarea")
        data[key] = "\n".join(f"Hallazgo {index}: " + "evidencia detallada " * 18 for index in range(30))
        content, _ = render_document(kind, context, common, data)
        presentation = Presentation(io.BytesIO(content))
        assert len(presentation.slides) > (10 if kind == "presentacion_proyecto" else 1)
        assert "Hallazgo 29" in texts(content, "pptx")
        assert all(shape.left + shape.width <= presentation.slide_width and shape.top + shape.height <= presentation.slide_height
                   for slide in presentation.slides for shape in slide.shapes)


def test_real_photo_is_embedded_with_caption_and_index(payload):
    from PIL import Image
    context, common = payload
    buffer = io.BytesIO()
    Image.new("RGB", (96, 48), (35, 80, 55)).save(buffer, format="PNG")
    content, _ = render_document("registro_evidencias", context, common, data_for("registro_evidencias"),
                                 images=[{"content": buffer.getvalue(), "caption": "Registro fotográfico de prueba", "fecha": date(2026, 10, 1), "autor": "Autor sintético"}])
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())
    text = texts(content, "docx")
    assert "Registro fotográfico de prueba" in text
    assert "01/10/2026" in text
    assert "Autor sintético" in text


def test_evidence_index_without_photos_marks_absence_and_never_invents_images(payload):
    context, common = payload
    content, _ = render_document("registro_evidencias", context, common, data_for("registro_evidencias"))
    assert "No se adjuntaron fotografías" in texts(content, "docx")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert not any(name.startswith("word/media/") for name in archive.namelist())


@pytest.mark.parametrize("changes", [{"kind": "tipo_inexistente"}, {"data": {"resultados": "no es una tabla"}},
    {"common": {"fecha_inicio": "fecha inválida"}}, {"common": {"presupuesto": [{"valor_planeado": "no es un valor"}]}},
    {"images": [{"content": b"no es una imagen", "caption": "Archivo corrupto"}]}])
def test_invalid_inputs_fail_with_actionable_error(payload, changes):
    context, common = payload
    args = {"kind": "registro_evidencias", "context": context, "common": common, "data": data_for("registro_evidencias")}
    if "data" in changes:
        args["kind"] = "informe_final"
        args["data"] = data_for("informe_final") | changes["data"]
    for key, value in changes.items():
        if key in {"data", "common"}:
            args[key] = args[key] | value
        else:
            args[key] = value
    with pytest.raises(ValueError) as error:
        render_document(**args)
    assert str(error.value)


def test_long_project_name_is_preserved_in_cover_body_below_a_separate_heading(payload):
    from pptx import Presentation
    context, common = payload
    context["nombre"] = "Fortalecimiento de los procesos de organización documental en entidades territoriales y públicas de la provincia de Vélez (referencia de validación)"
    content, _ = render_document("presentacion_proyecto", context, common, data_for("presentacion_proyecto"))
    slide = Presentation(io.BytesIO(content)).slides[0]
    heading, body = slide.shapes[0], slide.shapes[1]
    assert heading.text == "Presentación del proyecto"
    assert " ".join(body.text.split()).startswith(context["nombre"])
    assert heading.top + heading.height < body.top


def test_docx_table_headers_keep_with_the_first_record(payload):
    from docx import Document
    context, common = payload
    content, _ = render_document("acta_inicio", context, common, data_for("acta_inicio"))
    document = Document(io.BytesIO(content))
    assert document.tables
    assert all(paragraph.paragraph_format.keep_with_next for table in document.tables
               for cell in table.rows[0].cells for paragraph in cell.paragraphs)


def test_pptx_adapter_exposes_same_editable_artifact_contract(payload):
    from app.services.documentation_pptx import render_presentation, render_poster
    context, common = payload
    for kind, adapter in (("presentacion_proyecto", render_presentation), ("poster_producto", render_poster)):
        arguments = (context, common, data_for(kind)) + (([],) if kind == "poster_producto" else ())
        content, mime = adapter(*arguments)
        assert zipfile.is_zipfile(io.BytesIO(content))
        assert mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        assert "Borrador para revisión" in texts(content, "pptx")
        assert "SGPS-EJEMPLO" in texts(content, "pptx")


def test_poster_short_content_preserves_single_large_portrait_slide(payload):
    from pptx import Presentation
    context, common = payload
    content, _ = render_document("poster_producto", context, common, data_for("poster_producto"))
    poster = Presentation(io.BytesIO(content))
    assert len(poster.slides) == 1
    assert poster.slide_height > poster.slide_width
    assert abs(poster.slide_width / poster.slide_height - 0.75) < 0.001


def test_presentation_schedule_remains_an_editable_table(payload):
    from pptx import Presentation
    context, common = payload
    common["cronograma"] = [dict(common["cronograma"][0], actividad=f"Actividad {index}") for index in range(20)]
    content, _ = render_document("presentacion_proyecto", context, common, data_for("presentacion_proyecto"))
    presentation = Presentation(io.BytesIO(content))
    tables = [shape.table for slide in presentation.slides for shape in slide.shapes if shape.has_table]
    assert tables
    assert "Actividad 19" in "\n".join(cell.text for table in tables for row in table.rows for cell in row.cells)


def test_docx_title_has_no_decorative_border_and_table_rows_stay_together(payload):
    context, common = payload
    content, _ = render_document("acta_cierre", context, common, data_for("acta_cierre"))
    from docx import Document
    document = Document(io.BytesIO(content))
    assert document.styles["Title"].element.find("w:pPr/w:pBdr", W) is None
    assert all(row._tr.find("w:trPr/w:cantSplit", W) is not None for table in document.tables for row in table.rows)


def test_poster_includes_first_real_photo_in_single_slide(payload):
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    context, common = payload
    buffer = io.BytesIO()
    Image.new("RGB", (96, 48), (35, 80, 55)).save(buffer, format="PNG")
    content, _ = render_document("poster_producto", context, common, data_for("poster_producto"),
                                 images=[{"content": buffer.getvalue(), "caption": "Fotografía del producto"}])
    poster = Presentation(io.BytesIO(content))
    assert len(poster.slides) == 1
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in poster.slides[0].shapes)
    assert "Fotografía del producto" in texts(content, "pptx")


def test_draft_preserves_missing_optional_rows_as_pending_and_accepts_datetime(payload):
    from datetime import datetime
    context, common = payload
    context["objetivos_especificos"] = "Objetivo uno\nObjetivo dos"
    common["fecha_inicio"] = datetime(2026, 10, 1, 8, 30)
    data = data_for("acta_inicio")
    data["invitados"] = []
    content, _ = render_document("acta_inicio", context, common, data)
    text = texts(content, "docx")
    assert "Pendiente por diligenciar" in text
    assert "Objetivo uno" in text and "Objetivo dos" in text
    assert "01/10/2026" in text


@pytest.mark.parametrize("images", ["no es una lista", [None], [{"caption": "Sin bytes"}],
                                   [{"content": b""}], [{"content": b"a"}] * 51,
                                   [{"content": b"a" * (10 * 1024 * 1024 + 1)}]])
def test_photo_boundaries_reject_missing_bytes_and_excess_limits(payload, images):
    context, common = payload
    with pytest.raises(ValueError) as error:
        render_document("registro_evidencias", context, common, data_for("registro_evidencias"), images=images)
    assert str(error.value)


def test_poster_preserves_multiple_photos_on_extra_slides(payload):
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    context, common = payload
    buffer = io.BytesIO()
    Image.new("RGB", (48, 96), (35, 80, 55)).save(buffer, format="PNG")
    photos = [{"content": buffer.getvalue(), "caption": f"Fotografía {index}"} for index in range(2)]
    content, _ = render_document("poster_producto", context, common, data_for("poster_producto"), images=photos)
    poster = Presentation(io.BytesIO(content))
    assert len(poster.slides) == 2
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in poster.slides for shape in slide.shapes) == 2
    assert "Fotografía 1" in texts(content, "pptx")
    assert all(shape.left + shape.width <= poster.slide_width and shape.top + shape.height <= poster.slide_height
               for slide in poster.slides for shape in slide.shapes)


@pytest.mark.parametrize("changed", [{"common": []}, {"context": {}}, {"common": {"presupuesto": ["fila inválida"]}},
                                    {"common": {"presupuesto": [{"valor_planeado": True}]}},
                                    {"common": {"presupuesto": [{"valor_planeado": float('nan')}]}},
                                    {"data": {"tipo_cierre": "no confirmado"}}, {"common": {"responsable": 123}}])
def test_document_boundaries_reject_bad_shapes_numbers_and_closure_scope(payload, changed):
    context, common = payload
    args = {"kind": "acta_cierre", "context": context, "common": common, "data": data_for("acta_cierre")}
    for key, value in changed.items():
        args[key] = (args[key] | value) if isinstance(value, dict) and key != "context" else value
    with pytest.raises(ValueError) as error:
        render_document(**args)
    assert str(error.value)
