"""Los archivos Office conservan los datos institucionales y el aviso de validación."""

import io

import pytest

from app.services.documentation_catalog import DOCUMENT_DEFINITIONS
from app.services.documentation_renderers import render_document
from test_documentation_renderers import EXPECTED_KINDS, data_for, payload as payload, texts


@pytest.mark.parametrize("kind", sorted(EXPECTED_KINDS))
def test_every_artifact_keeps_source_notice_explicit(payload, kind):
    context, common = payload
    common["inconsistencias_fuente"] = "MARCADOR VALIDACIÓN: datos propuestos, pendientes de confirmación institucional."
    content, _ = render_document(kind, context, common, data_for(kind))
    artifact_text = " ".join(texts(content, DOCUMENT_DEFINITIONS[kind]["format"]).split())
    assert common["inconsistencias_fuente"] in artifact_text


def test_presentation_keeps_every_required_common_value_and_budget_as_editable_tables(payload):
    from pptx import Presentation

    context, common = payload
    common["regional"] = "REGIONAL PRESENTACIÓN"
    common["ciudad"] = "CIUDAD PRESENTACIÓN"
    common["equipo"][0].update(nombre="AUTORA PRESENTACIÓN", rol="ROL PRESENTACIÓN", actividades="ACTIVIDADES PRESENTACIÓN")
    common["presupuesto"][0].update(rubro="RUBRO PRESENTACIÓN", uso="USO PRESENTACIÓN")
    content, _ = render_document("presentacion_proyecto", context, common, data_for("presentacion_proyecto"))
    text = " ".join(texts(content, "pptx").split())
    for marker in ["REGIONAL PRESENTACIÓN", "CIUDAD PRESENTACIÓN", "AUTORA PRESENTACIÓN", "ROL PRESENTACIÓN",
                   "ACTIVIDADES PRESENTACIÓN", "RUBRO PRESENTACIÓN", "USO PRESENTACIÓN", "01/10/2026", "COP $ 10.996.585,25"]:
        assert marker in text
    tables = [shape.table for slide in Presentation(io.BytesIO(content)).slides for shape in slide.shapes if shape.has_table]
    assert any("RUBRO PRESENTACIÓN" in " ".join(cell.text for row in table.rows for cell in row.cells) for table in tables)


def test_poster_identifies_region_product_and_all_authors_without_contact_details(payload):
    context, common = payload
    common["regional"] = "REGIONAL PÓSTER"
    common["equipo"][0].update(nombre="AUTORA PÓSTER", identificacion="IDENTIFICACIÓN RESERVADA", correo_contacto="reservado@example.test")
    common["equipo"].append(dict(common["equipo"][0], nombre="SEGUNDA AUTORA PÓSTER"))
    context["producto"]["nombre"] = "PRODUCTO PÓSTER"
    content, _ = render_document("poster_producto", context, common, data_for("poster_producto"))
    text = " ".join(texts(content, "pptx").split())
    for marker in ["REGIONAL PÓSTER", "AUTORA PÓSTER", "SEGUNDA AUTORA PÓSTER", "PRODUCTO PÓSTER"]:
        assert marker in text
    assert "IDENTIFICACIÓN RESERVADA" not in text
    assert "reservado@example.test" not in text


def test_final_report_keeps_optional_institutional_context_and_planned_budget(payload):
    context, common = payload
    common["nivel_formacion"] = "NIVEL INFORME FINAL"
    common["programa_formacion"] = "PROGRAMA INFORME FINAL"
    common["competencia"] = "COMPETENCIA INFORME FINAL"
    common["presupuesto"][0]["uso"] = "USO PRESUPUESTAL INFORME FINAL"
    content, _ = render_document("informe_final", context, common, data_for("informe_final"))
    text = texts(content, "docx")
    for marker in ["NIVEL INFORME FINAL", "PROGRAMA INFORME FINAL", "COMPETENCIA INFORME FINAL", "USO PRESUPUESTAL INFORME FINAL"]:
        assert marker in text


def test_final_report_keeps_table_labels_and_values_on_the_same_page(payload):
    from docx import Document

    context, common = payload
    content, _ = render_document("informe_final", context, common, data_for("informe_final"))
    document = Document(io.BytesIO(content))
    for table in document.tables:
        for row in table.rows:
            assert row._tr.xpath("./w:trPr/w:cantSplit"), row.cells[0].text


def test_final_report_added_headings_are_black_and_keep_source_heading_style(payload):
    from docx import Document
    from docx.shared import RGBColor

    context, common = payload
    content, _ = render_document("informe_final", context, common, data_for("informe_final"))
    document = Document(io.BytesIO(content))
    start = next(index for index, paragraph in enumerate(document.paragraphs)
                 if paragraph.text == "Información complementaria del proyecto")
    headings = [paragraph for paragraph in document.paragraphs[start:]
                if paragraph.style.name.startswith("Heading")]
    assert len(headings) > 1
    for heading in headings:
        assert heading.paragraph_format.keep_with_next
        assert all(run.font.color.rgb == RGBColor(0, 0, 0) for run in heading.runs)
    source_heading = next(paragraph for paragraph in document.paragraphs
                          if paragraph.text.strip() == "Información general del proyecto")
    assert source_heading.style.name == "Heading 1"


def test_final_report_index_keeps_links_without_stale_template_page_numbers(payload):
    from docx import Document
    from docx.oxml.ns import qn

    context, common = payload
    content, _ = render_document("informe_final", context, common, data_for("informe_final"))
    document = Document(io.BytesIO(content))
    index = document.element.body.xpath("./w:sdt")[0]
    links = list(index.iter(qn("w:hyperlink")))
    assert len(links) == 16
    assert "Introducción" in " ".join(node.text or "" for node in index.iter(qn("w:t")))
    assert "14. Anexos" in " ".join(node.text or "" for node in index.iter(qn("w:t")))
    assert not list(index.iter(qn("w:instrText")))
    assert not list(index.iter(qn("w:fldChar")))
    assert all(not (node.text or "").isdigit() for node in index.iter(qn("w:t")))
