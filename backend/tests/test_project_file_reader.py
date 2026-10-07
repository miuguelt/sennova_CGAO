"""Lectura determinista y propuesta de campos verificables del expediente."""

import io
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from app.services.project_file_import import reader


def entry(path, content):
    return {"ruta": path, "nombre_archivo": path.rsplit("/", 1)[-1], "content": content}


def document_bytes(document):
    target = io.BytesIO()
    document.save(target)
    return target.getvalue()


@pytest.mark.parametrize("path,expected,period", [
    ("1ProyectoFomulado/formato.docx", "formulacion_proyecto", None),
    ("1ProyectoFomulado/presentacion.pptx", "presentacion_proyecto", None),
    ("2ActadeInicio/acta.docx", "acta_inicio", None),
    ("3Productos/2PosteryEventos/resultado.pptx", "poster_producto", None),
    ("5ActaCierre/informe_final.docx", "informe_final", None),
    ("4InformesBimensuales/Bimestre_2/informe.txt", "informe_bimensual", 2),
    ("6EvidenciasFotograficas/video.mp4", "evidencia_video", None),
    ("7Borradoresyvarios/nota.txt", "borrador_varios", None),
    ("anexos/nota.txt", "documento_apoyo", None),
])
def test_classification_uses_stage_and_specific_names(path, expected, period):
    result = reader.analyze_file(entry(path, b"Texto"))
    assert (result["tipo"], result["periodo_bimestre"]) == (expected, period)


def test_docx_reads_project_fields_labels_sections_and_tables():
    document = Document()
    table = document.add_table(rows=0, cols=2)
    for label, value in [("Título del Proyecto", "Estudio de cultivo"), ("Ciudad", "Vélez"), ("Código SGPS", "SGPS-20")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    document.add_heading("Metodología", 2)
    document.add_paragraph("Comparar parcelas durante seis meses.")
    document.add_heading("Equipo", 2)
    table = document.add_table(rows=1, cols=3)
    for cell, label in zip(table.rows[0].cells, ["Nombre", "Rol", "Responsabilidad"]):
        cell.text = label
    for cell, value in zip(table.add_row().cells, ["Investigador de prueba", "Líder", "Registrar resultados"]):
        cell.text = value
    result = reader.analyze_file(entry("1ProyectoFomulado/formato.docx", document_bytes(document)))
    assert result["propuesta"]["proyecto"]["nombre"] == "Estudio de cultivo"
    assert result["propuesta"]["proyecto"]["codigo_sgps"] == "SGPS-20"
    assert result["propuesta"]["comunes"]["ciudad"] == "Vélez"
    assert result["propuesta"]["borrador"]["metodologia"] == "Comparar parcelas durante seis meses."
    assert result["propuesta"]["comunes"]["equipo"][0]["nombre"] == "Investigador de prueba"
    assert "Registrar resultados" in result["texto_extraido"]


def test_text_csv_and_xlsx_extract_known_values_and_leave_unknowns_alone():
    for name, content in [("datos.txt", "Ciudad: Vélez\nRegional: Santander\nDato desconocido: no inventar".encode()),
                          ("datos.csv", "Ciudad;Vélez\nRegional;Santander".encode())]:
        result = reader.analyze_file(entry(name, content))
        assert result["propuesta"]["comunes"] == {"ciudad": "Vélez", "regional": "Santander"}
        assert result["propuesta"]["borrador"] == {}
    book = Workbook()
    book.active.append(["Ciudad", "Vélez"])
    book.active.append(["Regional", "Santander"])
    book.active.append(["Código CAP", "=1+1"])
    result = reader.analyze_file(entry("datos.xlsx", document_bytes(book)))
    assert result["propuesta"]["comunes"] == {"ciudad": "Vélez", "regional": "Santander"}
    assert any("fórmula" in warning for warning in result["advertencias"])


def test_pptx_reads_labels_and_content_without_inventing_fields():
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "Metodología"
    slide.placeholders[1].text = "Medir los cambios por parcela."
    result = reader.analyze_file(entry("presentacion.pptx", document_bytes(slides)))
    assert result["tipo"] == "presentacion_proyecto"
    assert result["propuesta"]["borrador"]["metodologia"] == "Medir los cambios por parcela."


@pytest.mark.parametrize("filename,content", [("a.pdf", b"%PDF-1.7"), ("a.png", b"PNG"), ("a.doc", b"documento antiguo"), ("a.mp4", b"video")])
def test_unsupported_extraction_keeps_original_with_clear_warning(filename, content):
    source = entry(filename, content)
    result = reader.analyze_file(source)
    assert result["texto_extraido"] == ""
    assert result["propuesta"] == {"comunes": {}, "borrador": {}, "proyecto": {}}
    assert any("original" in warning for warning in result["advertencias"])
    assert source["content"] == content


def test_invalid_dates_numbers_and_conflicting_labels_are_not_proposed():
    result = reader.analyze_file(entry("2ActadeInicio/acta.txt", (
        "Fecha de inicio del proyecto: ayer\nPresupuesto total: muchos\nCiudad: Vélez\nCiudad: Bogotá\n"
    ).encode()))
    assert "fecha_inicio" not in result["propuesta"]["comunes"]
    assert "presupuesto_total" not in result["propuesta"]["proyecto"]
    assert "ciudad" not in result["propuesta"]["comunes"]
    assert len(result["advertencias"]) >= 3


def test_explicit_kind_and_missing_bimester_are_validated():
    result = reader.analyze_file(entry("dato.txt", b"Ciudad: San Gil"), "informe_bimensual")
    assert result["tipo"] == "informe_bimensual"
    assert result["periodo_bimestre"] is None
    assert any("bimestre" in warning for warning in result["advertencias"])
    with pytest.raises(ValueError, match="tipo"):
        reader.analyze_file(entry("dato.txt", b"a"), "inventado")


def test_text_and_proposal_truncation_is_explicit(monkeypatch):
    monkeypatch.setattr(reader, "MAX_EXTRACTED_TEXT", 50)
    text = "Metodología: " + "a" * 21000
    result = reader.analyze_file(entry("1ProyectoFomulado/datos.txt", text.encode()))
    assert len(result["texto_extraido"]) == 50
    assert len(result["propuesta"]["borrador"]["metodologia"]) == 20000
    assert any("texto extraído" in warning for warning in result["advertencias"])
    assert any("Metodología" in warning and "recort" in warning for warning in result["advertencias"])


def test_damaged_or_hostile_office_produces_warning_without_guessing():
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w") as package:
        package.writestr("word/document.xml", '<!DOCTYPE test [<!ENTITY a "boom">]><doc/>')
    for content in (b"PK corrupto", target.getvalue()):
        result = reader.analyze_file(entry("1ProyectoFomulado/a.docx", content))
        assert result["propuesta"] == {"comunes": {}, "borrador": {}, "proyecto": {}}
        assert result["advertencias"]


def test_cap_fallback_does_not_restore_a_conflicting_title():
    document = Document()
    table = document.add_table(rows=0, cols=2)
    for value in ("Primer título", "Segundo título"):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = "Título del Proyecto", value
    result = reader.analyze_file(entry("1ProyectoFomulado/formato.docx", document_bytes(document)))
    assert "nombre" not in result["propuesta"]["proyecto"]
    assert any("valores diferentes" in warning for warning in result["advertencias"])


def test_cap_results_report_row_and_cell_truncation():
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "8. Resultados esperados\n" + "\n".join(f"Resultado {index}: " + "x" * 2100 for index in range(82))
    result = reader.analyze_file(entry("1ProyectoFomulado/formato.docx", document_bytes(document)))
    rows = result["propuesta"]["borrador"]["resultados_esperados"]
    assert len(rows) == 80 and len(rows[0]["resultado"]) == 2000
    assert any("80 filas" in warning for warning in result["advertencias"])
    assert any("2000 caracteres" in warning for warning in result["advertencias"])


def test_office_limits_and_invalid_xml_are_reported(monkeypatch):
    document = Document()
    document.add_paragraph("Ciudad: Vélez")
    monkeypatch.setattr(reader, "MAX_OFFICE_EXPANDED", 1)
    result = reader.analyze_file(entry("acta.docx", document_bytes(document)))
    assert "tamaño interno" in result["advertencias"][0]
    assert result["texto_extraido"] == ""


def test_row_extraction_currency_dates_selects_and_missing_fields():
    book = Workbook()
    sheet = book.active
    sheet.title = "Personal vinculado"
    sheet.append(["Nombre", "Rol", "Actividades a liderar"])
    sheet.append(["Persona de prueba", "Líder", "Medir parcelas"])
    metadata = book.create_sheet("Datos")
    metadata.append(["Presupuesto total", "COP $ 1.234.567,50"])
    metadata.append(["Duración en meses", 12])
    metadata.append(["Fecha de inicio del proyecto", "01/10/2026"])
    metadata.append(["Objetivos específicos", "Medir parcelas\nComparar resultados"])
    metadata.append(["Alcance del cierre", "Cierre final"])
    result = reader.analyze_file(entry("5ActaCierre/acta.xlsx", document_bytes(book)))
    assert result["propuesta"]["comunes"]["equipo"] == [{"nombre": "Persona de prueba", "rol": "Líder", "actividades": "Medir parcelas"}]
    assert result["propuesta"]["comunes"]["fecha_inicio"] == "2026-10-01"
    assert result["propuesta"]["proyecto"]["vigencia"] == 12
    assert result["propuesta"]["proyecto"]["presupuesto_total"] == "1234567.50"
    assert result["propuesta"]["proyecto"]["objetivos_especificos"] == ["Medir parcelas", "Comparar resultados"]
    assert result["propuesta"]["borrador"]["tipo_cierre"] == "final"


def test_csv_header_uniquely_identifies_a_catalog_table():
    source = entry("equipo.csv", "Nombre,Rol,Actividades a liderar\nInvestigador,Líder,Medir parcelas".encode())
    result = reader.analyze_file(source)
    assert result["propuesta"]["comunes"]["equipo"][0]["actividades"] == "Medir parcelas"


@pytest.mark.parametrize("filename", ["nota.txt", "audio.mp3", "datos.json"])
def test_non_visual_support_does_not_complete_photo_stage(filename):
    result = reader.analyze_file(entry("6EvidenciasFotograficas/" + filename, b"{}"))
    assert result["tipo"] == "documento_apoyo"
    assert any("visual" in warning for warning in result["advertencias"])


def test_bimester_alias_and_json_manifest_do_not_change_project_data():
    result = reader.analyze_file(entry("Bimestre_2.txt", b"Texto"), "informe_bimestral")
    assert result["tipo"] == "informe_bimensual" and result["periodo_bimestre"] == 2
    result = reader.analyze_file(entry("expediente.json", b'{"ciudad": "Velez"}'))
    assert result["texto_extraido"] == '{"ciudad": "Velez"}'
    assert result["propuesta"] == {"comunes": {}, "borrador": {}, "proyecto": {}}


def test_product_report_folder_is_not_mistaken_for_final_project_report():
    result = reader.analyze_file(entry("3Productos/1InformeFinal/resultado.txt", b"Contenido"))
    assert result["tipo"] == "producto_resultado"


def test_explicit_visual_type_still_requires_visual_file():
    with pytest.raises(ValueError, match="imagen"):
        reader.analyze_file(entry("a.txt", b"Texto"), "evidencia_fotografica")


def test_bimester_is_read_from_explicit_content_and_conflicts_stay_pending():
    result = reader.analyze_file(entry("4InformesBimensuales/informe.txt", "Período bimestral 3\nCiudad: Vélez".encode()))
    assert result["periodo_bimestre"] == 3
    assert not any("No se identificó" in warning for warning in result["advertencias"])
    conflict = reader.analyze_file(entry("4InformesBimensuales/Bimestre_2.txt", b"Periodo bimestral 3"))
    assert conflict["periodo_bimestre"] is None
    assert any("bimestre" in warning for warning in conflict["advertencias"])


def test_multiline_labeled_narrative_keeps_all_paragraphs():
    result = reader.analyze_file(entry("1ProyectoFomulado/formulacion.txt", b"Metodologia: Medir parcelas.\nComparar resultados.\nCiudad: Velez"))
    assert result["propuesta"]["borrador"]["metodologia"] == "Medir parcelas.\nComparar resultados."
