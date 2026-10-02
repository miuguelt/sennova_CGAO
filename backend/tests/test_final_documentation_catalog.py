"""Contenido del informe final basado en una plantilla adaptable, sin datos reales."""

import io
import os
import secrets
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from fastapi import HTTPException
from docx import Document

os.environ.setdefault("DATABASE_URL", "sqlite://")
os.environ.setdefault("DEBUG", "true")
os.environ.setdefault("JWT_SECRET", secrets.token_urlsafe(32))

from app.services.documentation_catalog import DOCUMENT_DEFINITIONS, RESULT_FIELDS
from app.services.documentation_renderers import render_document
from app.services.documentation_final_template import _development_lines, _result_lines
from app.services.documentation_state import generation_pending
from app.services.documentation_validation import validate_fields

NARRATIVE_KEYS = {
    "introduccion", "planteamiento_problema", "estado_arte_tecnica", "desarrollo_proyecto",
    "viabilidad_tecnica", "viabilidad_operativa", "viabilidad_economica", "viabilidad_normativa",
    "viabilidad_mercado", "propiedad_intelectual_transferencia", "impacto_proyecto", "anexos",
}
REQUIRED_METADATA = {"autor_informe", "fecha_entrega", "clasificacion_informacion"}
OPTIONAL_METADATA = {"codigo_idea", "experto_proyecto", "linea_tecnologica", "trl_inicial", "trl_alcanzado", "tecnoparque"}


def test_final_catalog_adds_specific_guidance_without_changing_other_document_types():
    fields = {field["key"]: field for field in DOCUMENT_DEFINITIONS["informe_final"]["fields"]}
    assert NARRATIVE_KEYS | REQUIRED_METADATA | OPTIONAL_METADATA <= fields.keys()
    assert all(fields[key]["required"] for key in NARRATIVE_KEYS | REQUIRED_METADATA)
    assert all(not fields[key]["required"] and "cuando aplique" in fields[key]["help"].lower()
               for key in OPTIONAL_METADATA)
    assert all("No aplica" in fields[key]["help"] for key in NARRATIVE_KEYS)
    assert all(field in DOCUMENT_DEFINITIONS["informe_final"]["fields"] for field in RESULT_FIELDS)
    assert {"cumplimiento_objetivos", "balance_final"} <= fields.keys()
    assert len(fields) == len(DOCUMENT_DEFINITIONS["informe_final"]["fields"])
    assert DOCUMENT_DEFINITIONS["informe_final"]["format"] == "docx"
    assert DOCUMENT_DEFINITIONS["informe_final"]["folder"] == "5ActaCierre"
    final_only = (NARRATIVE_KEYS - {"introduccion", "planteamiento_problema"}) | REQUIRED_METADATA | OPTIONAL_METADATA
    for kind, definition in DOCUMENT_DEFINITIONS.items():
        if kind != "informe_final":
            assert final_only.isdisjoint(field["key"] for field in definition["fields"])


def test_final_required_metadata_and_narratives_are_pending_but_technoparque_fields_are_optional():
    project = SimpleNamespace(objetivo_general="Organizar documentos de prueba", objetivos_especificos=["Medir el resultado"])
    slot = {"tipo": "informe_final"}
    pending = {item["campo"] for item in generation_pending(project, slot, {}, {})}
    assert NARRATIVE_KEYS | REQUIRED_METADATA <= pending
    assert OPTIONAL_METADATA.isdisjoint(pending)
    fields = DOCUMENT_DEFINITIONS["informe_final"]["fields"]
    normalized = validate_fields({key: "No aplica: el alcance documental no contempla esta actividad."
                                  for key in NARRATIVE_KEYS}, fields)
    assert set(normalized) == NARRATIVE_KEYS
    assert NARRATIVE_KEYS.isdisjoint({item["campo"] for item in generation_pending(project, slot, {}, normalized)})


@pytest.mark.parametrize("value", [0, 9])
def test_final_trl_accepts_inclusive_zero_to_nine(value):
    normalized = validate_fields({"trl_inicial": value, "trl_alcanzado": value}, DOCUMENT_DEFINITIONS["informe_final"]["fields"])
    assert normalized == {"trl_inicial": str(value), "trl_alcanzado": str(value)}


@pytest.mark.parametrize("key,value", [
    ("trl_inicial", -1), ("trl_alcanzado", 10), ("trl_inicial", True),
    ("clasificacion_informacion", "interna"), ("fecha_entrega", "2026-02-30"),
])
def test_final_metadata_rejects_invalid_ranges_classification_and_delivery_date(key, value):
    with pytest.raises(HTTPException) as error:
        validate_fields({key: value}, DOCUMENT_DEFINITIONS["informe_final"]["fields"])
    assert error.value.status_code == 422
    assert "no reconocidos" not in error.value.detail


@pytest.mark.parametrize("classification", ["publica", "publica_clasificada", "publica_reservada"])
def test_final_classification_and_delivery_date_preserve_explicit_values(classification):
    values = {"clasificacion_informacion": classification, "fecha_entrega": "2026-10-01"}
    assert validate_fields(values, DOCUMENT_DEFINITIONS["informe_final"]["fields"]) == values


def test_final_renderer_emits_every_new_field_label_and_explicit_value_in_real_docx():
    fields = {field["key"]: field for field in DOCUMENT_DEFINITIONS["informe_final"]["fields"]}
    added = NARRATIVE_KEYS | REQUIRED_METADATA | OPTIONAL_METADATA
    data = {key: "Dato sintético para " + key for key in added}
    data.update(fecha_entrega="2026-10-01", clasificacion_informacion="publica_clasificada", trl_inicial="0", trl_alcanzado="9")
    context = {"nombre": "Proyecto sintético de informe final", "objetivo_general": "Organizar archivos",
               "objetivos_especificos": ["Clasificar documentos"]}
    content, mime = render_document("informe_final", context, {}, data)
    assert mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    with ZipFile(io.BytesIO(content)) as document:
        rendered = Document(io.BytesIO(content))
        text = "\n".join([paragraph.text for paragraph in rendered.paragraphs]
                         + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells])
        for key in added:
            label = {
                "autor_informe": "Talento que realiza el informe",
                "fecha_entrega": "Fecha de entrega",
                "codigo_idea": "Código de la Idea",
                "clasificacion_informacion": "CLASIFICACIÓN DE LA INFORMACIÓN",
                "tecnoparque": "TecnoParque en el que fue desarrollado el Proyecto de Base Tecnológica",
            }.get(key, fields[key]["label"])
            assert label in text
            expected = "Pública Clasificada" if key == "clasificacion_informacion" else "01/10/2026" if key == "fecha_entrega" else data[key]
            assert expected in text
        assert context["nombre"] in text
        assert "GCDTP-F-023" not in text


def test_final_report_sections_preserve_source_labels_and_values():
    fields = {field["key"]: field for field in DOCUMENT_DEFINITIONS["informe_final"]["fields"]}
    data = {
        "antecedentes": "Contexto documentado",
        "desarrollo_proyecto": "Implementación reportada",
        "discusion": "Interpretación del resultado",
        "fortalezas": "Equipo vinculado",
        "dificultades": "Pendiente registrado",
        "cumplimiento_objetivos": "Avance frente a objetivos",
    }

    development = _development_lines(fields, data)
    results = _result_lines(fields, data)

    assert "Antecedentes: Contexto documentado" in development
    assert "Desarrollo del proyecto: Implementación reportada" in development
    assert "Discusión de resultados: Interpretación del resultado" in results
    assert "Dificultades y limitaciones: Pendiente registrado" in results
