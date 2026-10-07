"""Importación completa con archivos y base de datos reales aislados."""

import io
import json
import uuid
from zipfile import ZipFile

import pytest
from docx import Document

from test_project_evidence_file import evidence_context  # noqa: F401


def docx_content():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Título del Proyecto"
    table.cell(0, 1).text = "Nombre leído del archivo"
    document.add_paragraph("2. Introducción")
    document.add_paragraph("La investigación organiza los documentos del proyecto.")
    document.add_paragraph("3. Planteamiento del problema")
    document.add_paragraph("Los registros requieren una organización verificable.")
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


def zip_content():
    target = io.BytesIO()
    with ZipFile(target, "w") as archive:
        archive.writestr("1ProyectoFomulado/formulacion.docx", docx_content())
        archive.writestr("6EvidenciasFotograficas/visita/notas.txt", "Visita de campo para revisar el avance.")
        archive.writestr("7Borradoresyvarios/vacia/", "")
    return target.getvalue()


def analyze(ctx, files=None, **data):
    return ctx[-1].post(f"/proyectos/{ctx[2].id}/expediente/analizar-archivos",
                        files=files or [("files", ("proyecto.zip", zip_content(), "application/zip"))], data=data)


def confirm(ctx, preview, files=None, **data):
    selection = [{"ruta": row["ruta"], "sha256": row["sha256"],
                  "periodo_bimestre": row["periodo_bimestre"], "importar_datos": True}
                 for row in preview["archivos"]]
    return ctx[-1].post(f"/proyectos/{ctx[2].id}/expediente/importar-archivos",
                        files=files or [("files", ("proyecto.zip", zip_content(), "application/zip"))],
                        data={**data, "seleccion": json.dumps(selection)})


def test_preview_does_not_write_and_commit_preserves_folders_data_and_originals(evidence_context):
    from app.models import Documento
    from app.documentation_models import ProjectDocumentDraft
    ctx = evidence_context
    files = [("files", ("proyecto.zip", zip_content(), "application/zip"))]
    preview = analyze(ctx, files)
    assert preview.status_code == 200, preview.text
    assert ctx[0].query(Documento).count() == 0
    assert not (ctx[3] / "documentos").exists()
    assert len(preview.json()["archivos"]) == 2
    assert "2ActadeInicio" in preview.json()["carpetas_faltantes"]
    result = confirm(ctx, preview.json(), files)
    assert result.status_code == 201, result.text
    assert result.json()["archivos_importados"] == 2
    assert result.json()["campos_registrados"] >= 1
    assert ctx[2].nombre == "Proyecto de organización documental"
    assert any("nombre" in warning.lower() for warning in result.json()["advertencias"])
    draft = ctx[0].query(ProjectDocumentDraft).filter_by(proyecto_id=ctx[2].id, clave="formulacion_proyecto").one()
    assert "organiza los documentos" in draft.datos["introduccion"]
    exported = ctx[-1].get(f"/proyectos/{ctx[2].id}/expediente/descargar")
    with ZipFile(io.BytesIO(exported.content)) as archive:
        assert "7Borradoresyvarios/vacia/" in archive.namelist()
        assert archive.read("6EvidenciasFotograficas/visita/notas.txt") == b"Visita de campo para revisar el avance."
        original = ZipFile(io.BytesIO(files[0][1][1])).read("1ProyectoFomulado/formulacion.docx")
        assert archive.read("1ProyectoFomulado/formulacion.docx") == original
    doc = ctx[0].query(Documento).filter_by(tipo="formulacion_proyecto").one()
    assert ctx[-1].get(f"/documentos/{doc.id}").status_code == 200


def test_repeated_upload_is_idempotent_and_download_names_do_not_collide(evidence_context):
    ctx = evidence_context
    files = [("files", ("notas.txt", b"Primera version", "text/plain"))]
    preview = analyze(ctx, files, carpeta="7Borradoresyvarios").json()
    first = confirm(ctx, preview, files, carpeta="7Borradoresyvarios")
    assert first.status_code == 201, first.text
    repeated = confirm(ctx, preview, files, carpeta="7Borradoresyvarios")
    assert repeated.json()["archivos_importados"] == 0
    assert repeated.json()["archivos_omitidos"] == 1
    newer = [("files", ("notas.txt", b"Segunda version", "text/plain"))]
    assert confirm(ctx, analyze(ctx, newer, carpeta="7Borradoresyvarios").json(), newer,
                   carpeta="7Borradoresyvarios").status_code == 201
    with ZipFile(io.BytesIO(ctx[-1].get(f"/proyectos/{ctx[2].id}/expediente/descargar").content)) as archive:
        names = [name for name in archive.namelist() if name.endswith("notas.txt")]
        assert len(names) == 2
        assert {archive.read(name) for name in names} == {b"Primera version", b"Segunda version"}


def test_hash_mismatch_rejects_without_writes(evidence_context):
    from app.models import Documento
    ctx = evidence_context
    files = [("files", ("notas.txt", b"Contenido", "text/plain"))]
    preview = analyze(ctx, files).json()
    preview["archivos"][0]["sha256"] = "0" * 64
    response = confirm(ctx, preview, files)
    assert response.status_code == 409
    assert ctx[0].query(Documento).count() == 0


@pytest.mark.parametrize("selection", ["no es JSON", "{}", "[]", '[{"ruta":"x"}]'])
def test_invalid_selection_rejects(evidence_context, selection):
    ctx = evidence_context
    result = ctx[-1].post(f"/proyectos/{ctx[2].id}/expediente/importar-archivos",
                          files=[("files", ("notas.txt", b"Contenido", "text/plain"))],
                          data={"seleccion": selection})
    assert result.status_code == 422


def test_report_requires_a_valid_period(evidence_context):
    ctx = evidence_context
    files = [("files", ("informe.docx", docx_content(), "application/octet-stream"))]
    preview = analyze(ctx, files, tipo="informe_bimensual").json()
    assert confirm(ctx, preview, files, tipo="informe_bimensual").status_code == 422
    preview["archivos"][0]["periodo_bimestre"] = 9
    assert confirm(ctx, preview, files, tipo="informe_bimensual").status_code == 422
    preview["archivos"][0]["periodo_bimestre"] = 2
    assert confirm(ctx, preview, files, tipo="informe_bimensual").status_code == 201


def test_permissions_and_missing_project(evidence_context):
    ctx = evidence_context
    ctx[1].rol = "aprendiz"
    assert analyze(ctx).status_code == 403
    ctx[1].rol = "investigador"
    result = ctx[-1].post(f"/proyectos/{uuid.uuid4()}/expediente/analizar-archivos",
                          files=[("files", ("notas.txt", b"Contenido", "text/plain"))])
    assert result.status_code == 404


def test_persistence_failure_rolls_back_files_and_rows(evidence_context, monkeypatch):
    from app.models import Documento
    ctx = evidence_context
    files = [("files", ("notas.txt", b"Contenido", "text/plain"))]
    preview = analyze(ctx, files).json()
    def fail_commit():
        raise OSError("Fallo de almacenamiento de prueba")
    monkeypatch.setattr(ctx[0], "commit", fail_commit)
    result = confirm(ctx, preview, files)
    assert result.status_code == 500
    assert ctx[0].query(Documento).count() == 0
    assert list((ctx[3] / "documentos").glob("*")) == []
