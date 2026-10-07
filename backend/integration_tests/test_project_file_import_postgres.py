"""Carga, lectura, regeneración y descarga en un esquema PostgreSQL aislado."""

import io
import json
from zipfile import ZipFile

import pytest
from docx import Document

from test_project_evidence_postgres import postgres_context  # noqa: F401
from app.routers import project_file_import, project_files, project_documentation
from app.services.project_file_import.models import ProjectFileBatch, ProjectImportedFile
from app.services.project_file_import.schema import downgrade_file_import_schema, upgrade_file_import_schema
from app.documentation_models import ProjectDocumentDraft
from app.models import Documento


def test_postgres_import_round_trip_and_migration_protection(postgres_context):
    engine, db, _, project, _, client = postgres_context
    for router in (project_file_import.router, project_files.router, project_documentation.router):
        client.app.include_router(router)
    document = Document()
    document.add_paragraph("2. Introducción")
    document.add_paragraph("El equipo organiza las fuentes del proyecto.")
    document.add_paragraph("3. Planteamiento del problema")
    document.add_paragraph("Se requiere clasificar los registros.")
    word = io.BytesIO()
    document.save(word)
    package = io.BytesIO()
    with ZipFile(package, "w") as archive:
        archive.writestr("1ProyectoFomulado/formulacion.docx", word.getvalue())
        archive.writestr("7Borradoresyvarios/notas/", "")
    files = [("files", ("proyecto.zip", package.getvalue(), "application/zip"))]
    base = f"/proyectos/{project.id}/expediente"
    preview = client.post(base + "/analizar-archivos", files=files)
    assert preview.status_code == 200
    assert db.query(Documento).count() == 0
    selected = [{"ruta": item["ruta"], "sha256": item["sha256"]} for item in preview.json()["archivos"]]
    saved = client.post(base + "/importar-archivos", files=files, data={"seleccion": json.dumps(selected)})
    assert saved.status_code == 201 and saved.json()["archivos_importados"] == 1
    db.expire_all()
    source = db.query(ProjectImportedFile).one()
    draft = db.query(ProjectDocumentDraft).filter_by(clave="formulacion_proyecto").one()
    assert draft.datos["introduccion"] == "El equipo organiza las fuentes del proyecto."
    assert source.propuesta["borrador"]["introduccion"] == draft.datos["introduccion"]
    duplicate = client.post(base + "/importar-archivos", files=files, data={"seleccion": json.dumps(selected)})
    assert duplicate.json()["archivos_omitidos"] == 1
    download = client.get(base + "/descargar")
    assert download.status_code == 200
    with ZipFile(io.BytesIO(download.content)) as archive:
        assert archive.read("1ProyectoFomulado/formulacion.docx") == word.getvalue()
        assert "7Borradoresyvarios/notas/" in archive.namelist()
    db.rollback()
    assert upgrade_file_import_schema(engine) == []
    with pytest.raises(ValueError, match="existen carpetas"):
        downgrade_file_import_schema(engine)
    assert db.query(ProjectFileBatch).count() == 1


def test_postgres_empty_migration_can_be_reversed_and_restored(postgres_context):
    engine, db, *_ = postgres_context
    db.rollback()
    assert downgrade_file_import_schema(engine) == ["project_imported_files", "project_file_batches"]
    assert upgrade_file_import_schema(engine) == ["project_file_batches", "project_imported_files"]
    assert upgrade_file_import_schema(engine) == []
    assert db.query(ProjectImportedFile).count() == 0
