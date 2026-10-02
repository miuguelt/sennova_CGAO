"""Generación DOCX y contrato público de borradores Office desde datos estructurados."""

import io

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from app.services.documentation_catalog import COMMON_FIELDS, DOCUMENT_DEFINITIONS
from app.services.documentation_final_template import render_final_report
from app.services.documentation_pptx import render_poster, render_presentation
from app.services.documentation_rendering_values import PENDING, format_value, photo_records, validate_fields

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _project_rows(context, common):
    """Conserva la información institucional registrada y el contexto del proyecto."""
    rows = [("Nombre del proyecto", context["nombre"]), ("Código SGPS", str(context.get("codigo_sgps") or PENDING))]
    for field in COMMON_FIELDS:
        if field["type"] not in {"rows", "textarea"}:
            rows.append((field["label"], format_value(common.get(field["key"]), field)))
    rows.extend([
        ("Grupo de investigación", str(context.get("grupo") or PENDING)),
        ("Semillero", str(context.get("semillero") or PENDING)),
        ("Duración en meses", str(context.get("vigencia") or PENDING)),
        ("Presupuesto total", format_value(context.get("presupuesto_total"), {"label": "Presupuesto total", "type": "number", "unit": "COP"})),
    ])
    if context.get("producto"):
        product = context["producto"]
        rows.extend([("Producto vinculado", str(product.get("nombre") or PENDING)), ("Tipología del producto", str(product.get("tipo") or PENDING))])
    return rows



def _style_docx(document):
    """Define legibilidad, jerarquía y márgenes para tablas que crecen."""
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.59), Cm(27.94)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    section.header_distance = section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = False
    document.settings.odd_and_even_pages_header_footer = False
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(10 if name == "Normal" else 20 if name == "Title" else 13 if name == "Heading 1" else 11)
        style.paragraph_format.space_after = Pt(7)
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
    document.styles["Normal"].paragraph_format.line_spacing = 1.15



def _table(document, headers, rows, *, signature=False):
    """Genera una tabla de filas expansibles con encabezados repetidos y firmas vacías."""
    table = document.add_table(rows=1, cols=len(headers) + int(signature))
    labels = list(headers) + (["Firma"] if signature else [])
    for cell, label in zip(table.rows[0].cells, labels):
        cell.text = label
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
        for run in cell.paragraphs[0].runs:
            run.bold = True
    header_repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header_repeat)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
        if signature:
            cells[-1].text = ""
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{edge}")
        for attribute, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            border.set(qn(f"w:{attribute}"), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge in ["top", "left", "bottom", "right"]:
        margin = OxmlElement(f"w:{edge}")
        margin.set(qn("w:w"), "100")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    table._tbl.tblPr.append(margins)
    for index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EEF3")
                cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                for run in paragraph.runs:
                    run.font.size = Pt(8 if len(labels) >= 7 else 9)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)
    return table



def _write_field(document, field, values):
    """Escribe un campo narrativo o una tabla sin truncar párrafos ni filas."""
    document.add_heading(field["label"], level=2)
    if field["type"] == "rows":
        columns = field["columns"]
        rows = [[format_value(row.get(column["key"]), column) for column in columns] for row in (values.get(field["key"]) or [])]
        if not rows:
            rows = [[PENDING] + [""] * (len(columns) - 1)]
        _table(document, [column["label"] for column in columns], rows, signature=field["key"] == "asistentes")
    else:
        text = format_value(values.get(field["key"]), field)
        for paragraph in text.splitlines():
            if paragraph.strip():
                document.add_paragraph(paragraph)



def _write_objectives(document, context):
    """Reutiliza objetivos del proyecto sin duplicarlos en el formulario."""
    document.add_heading("Objetivo general", level=2)
    document.add_paragraph(str(context.get("objetivo_general") or PENDING))
    document.add_heading("Objetivos específicos", level=2)
    objectives = context.get("objetivos_especificos") or []
    if isinstance(objectives, str):
        objectives = objectives.splitlines()
    for objective in objectives or [PENDING]:
        document.add_paragraph(str(objective), style="List Bullet")



def _render_docx(kind, context, common, data, photos):
    """Construye secciones de las actas y los informes con soportes editables."""
    definition = DOCUMENT_DEFINITIONS[kind]
    document = Document()
    _style_docx(document)
    document.add_paragraph(definition["title"], style="Title")
    document.add_paragraph(context["nombre"], style="Subtitle")
    status = f"Borrador para revisión · Versión {context.get('version', 1)}"
    document.add_paragraph(status)
    document.sections[0].header.paragraphs[0].text = status
    document.sections[0].footer.paragraphs[0].text = "Las firmas y las aprobaciones requieren revisión y registro independiente."
    if kind == "informe_bimensual":
        document.add_paragraph(f"Período bimestral {context.get('periodo_bimestre') or PENDING}")
    if kind in {"acta_inicio", "acta_cierre"}:
        document.add_heading("Datos de la reunión", level=1)
        metadata = [field for field in definition["fields"] if field["key"] in {"tipo_cierre", "periodo_desde", "periodo_hasta", "fecha_reunion", "hora_inicio", "hora_fin", "lugar"}]
        _table(document, ["Dato", "Información registrada"], [(field["label"], format_value(data.get(field["key"]), field)) for field in metadata])
        for field in definition["fields"]:
            if field["key"] in {"temas", "objetivo_reunion"}:
                _write_field(document, field, data)
    document.add_heading("Información general del proyecto", level=1)
    _table(document, ["Dato", "Información registrada"], _project_rows(context, common))
    _write_objectives(document, context)
    if kind in {"acta_inicio", "formulacion_proyecto"}:
        for field in COMMON_FIELDS:
            if field["type"] == "rows":
                _write_field(document, field, common)
    skipped = {"tipo_cierre", "periodo_desde", "periodo_hasta", "fecha_reunion", "hora_inicio", "hora_fin", "lugar", "temas", "objetivo_reunion"} if kind in {"acta_inicio", "acta_cierre"} else set()
    document.add_heading("Desarrollo y soportes", level=1)
    for field in definition["fields"]:
        if field["key"] not in skipped:
            _write_field(document, field, data)
    if common.get("inconsistencias_fuente"):
        _write_field(document, next(field for field in COMMON_FIELDS if field["key"] == "inconsistencias_fuente"), common)
    if kind == "registro_evidencias" or photos:
        document.add_heading("Registro fotográfico", level=1)
        if not photos:
            document.add_paragraph("No se adjuntaron fotografías. El índice debe vincular los soportes reales disponibles.")
        for content, caption in photos:
            with Image.open(io.BytesIO(content)) as image:
                width = min(Cm(15), int(Cm(16) * image.width / image.height))
            document.add_picture(io.BytesIO(content), width=width)
            document.paragraphs[-1].paragraph_format.keep_with_next = True
            document.add_paragraph(caption, style="Caption")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), DOCX_MIME



def render_document(kind, context: dict, common: dict, data: dict, *, images=None) -> tuple[bytes, str]:
    """Genera un borrador real sin consultas de base de datos ni acceso a archivos."""
    if kind not in DOCUMENT_DEFINITIONS:
        raise ValueError("El tipo de documento solicitado no está disponible")
    if not isinstance(context, dict) or not isinstance(context.get("nombre"), str) or not context["nombre"].strip():
        raise ValueError("El proyecto debe tener un nombre registrado")
    validate_fields(COMMON_FIELDS, common)
    validate_fields(DOCUMENT_DEFINITIONS[kind]["fields"], data)
    photos = photo_records(images)
    if kind == "poster_producto":
        return render_poster(context, common, data, photos)
    if DOCUMENT_DEFINITIONS[kind]["format"] == "pptx":
        return render_presentation(context, common, data)
    if kind == "informe_final":
        return render_final_report(context, common, data, photos)
    return _render_docx(kind, context, common, data, photos)

