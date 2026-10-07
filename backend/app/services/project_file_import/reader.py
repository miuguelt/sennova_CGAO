"""Lee fuentes en memoria y propone solo datos respaldados por rótulos conocidos."""

import csv
import io
import re
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime
from pathlib import PurePosixPath

from fastapi import HTTPException
from openpyxl import load_workbook

from app.services.documentation_catalog import COMMON_FIELDS, DOCUMENT_DEFINITIONS
from app.services.documentation_validation import normalize_value, validate_fields
from app.services.project_evidence_service import STAGES
from app.services.proyecto_import_service import GUIDED_SECTION_FIELDS, _extract_objectives, _find_labeled_value, _section_content
from app.services.project_file_import.archive import safe_path

MAX_EXTRACTED_TEXT = 100000
MAX_OFFICE_EXPANDED = 30 * 1024 * 1024
MAX_OFFICE_ENTRIES = 2048
MAX_SHEET_CELLS = 50000
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
PROJECT_FIELDS = (
    {"key": "nombre", "label": "Nombre del proyecto", "type": "text"},
    {"key": "codigo_sgps", "label": "Código SGPS", "type": "text"},
    {"key": "objetivo_general", "label": "Objetivo general", "type": "textarea"},
    {"key": "objetivos_especificos", "label": "Objetivos específicos", "type": "textarea"},
    {"key": "vigencia", "label": "Duración en meses", "type": "number", "min": 1, "max": 60},
    {"key": "presupuesto_total", "label": "Presupuesto total", "type": "number", "unit": "COP"},
)
_EMPTY = {"pendiente por diligenciar", "pendiente de completar", "sin informacion", "no aplica"}
_IMAGES = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".heic", ".heif"}
_VIDEOS = {".mp4", ".mov", ".avi", ".webm", ".mkv"}


def _normalize(text):
    plain = "".join(char for char in unicodedata.normalize("NFKD", str(text).casefold()) if not unicodedata.combining(char))
    plain = re.sub(r"^\s*(?:\d+[.)\-]\s*)+", "", plain)
    return re.sub(r"[^a-z0-9]+", " ", plain).strip()


def _classify(path, explicit):
    known = set().union(*(stage[4] for stage in STAGES))
    if explicit:
        if explicit not in known:
            raise ValueError("El tipo de documento seleccionado no es válido.")
        if explicit == "evidencia_fotografica" and PurePosixPath(path).suffix.lower() not in _IMAGES:
            raise ValueError("La evidencia fotográfica debe ser un archivo de imagen admitido.")
        if explicit == "evidencia_video" and PurePosixPath(path).suffix.lower() not in _VIDEOS:
            raise ValueError("La evidencia de video debe ser un archivo de video admitido.")
        return "informe_bimensual" if explicit == "informe_bimestral" else explicit
    text = _normalize(path)
    filename = _normalize(PurePosixPath(path).name)
    if "informe final" in filename or "informefinal" in filename:
        return "informe_final"
    if "poster" in text:
        return "poster_producto"
    if "presentacion" in text:
        return "presentacion_proyecto"
    for stage in STAGES:
        if any(_normalize(part) == _normalize(stage[1]) for part in PurePosixPath(path).parts[:-1]):
            if stage[0] == "evidencias":
                extension = PurePosixPath(path).suffix.lower()
                return "evidencia_video" if extension in _VIDEOS else "evidencia_fotografica" if extension in _IMAGES else "documento_apoyo"
            return stage[3]
    for kind in ("formulacion_proyecto", "acta_inicio", "acta_cierre", "informe_bimensual", "informe_bimestral"):
        if _normalize(kind) in text:
            return "informe_bimensual" if kind == "informe_bimestral" else kind
    return "documento_apoyo"


def _office_parts(content):
    """Verifica los límites del contenedor y prohíbe DTD antes de analizar XML."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as package:
            entries = package.infolist()
            if len(entries) > MAX_OFFICE_ENTRIES or sum(item.file_size for item in entries) > MAX_OFFICE_EXPANDED:
                raise ValueError("El documento Office supera el tamaño interno permitido.")
            parts = {}
            seen = set()
            for item in entries:
                name = safe_path(item.orig_filename.rstrip("/") if item.is_dir() else item.orig_filename)
                if name.casefold() in seen or item.flag_bits & 1:
                    raise ValueError("El documento Office contiene entradas duplicadas o cifradas.")
                seen.add(name.casefold())
                if name.endswith((".xml", ".rels")):
                    value = package.read(item)
                    check = value.replace(b"\x00", b"").upper()
                    if b"<!DOCTYPE" in check or b"<!ENTITY" in check:
                        raise ValueError("El documento Office contiene una estructura XML no permitida.")
                    parts[name] = ET.fromstring(value)
            return parts
    except (zipfile.BadZipFile, ET.ParseError, KeyError, OSError, RuntimeError, EOFError, NotImplementedError) as error:
        raise ValueError("El documento Office está dañado o no se puede leer.") from error


def _paragraph_text(node, namespace):
    pieces = []
    for child in node.iter():
        if child.tag == namespace + "t":
            pieces.append(child.text or "")
        elif child.tag in {namespace + "br", namespace + "cr"}:
            pieces.append("\n")
        elif child.tag == namespace + "tab":
            pieces.append(" ")
    return "".join(pieces).strip()


def _docx_blocks(parts):
    document = parts.get("word/document.xml")
    body = document.find(W + "body") if document is not None else None
    if body is None:
        raise ValueError("El DOCX no contiene un cuerpo de documento legible.")
    blocks = []
    for node in body:
        if node.tag == W + "p":
            style = node.find(W + "pPr/" + W + "pStyle")
            heading = style is not None and re.search(r"heading|titulo|título", style.get(W + "val", ""), re.I)
            blocks.append(("heading" if heading else "text", _paragraph_text(node, W)))
        elif node.tag == W + "tbl":
            rows = [["\n".join(_paragraph_text(paragraph, W) for paragraph in cell.findall(W + "p"))
                     for cell in row.findall(W + "tc")] for row in node.findall(W + "tr")]
            blocks.append(("table", rows))
    return blocks


def _pptx_blocks(parts):
    names = sorted((name for name in parts if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
                   key=lambda name: int(re.search(r"slide(\d+)\.xml", name).group(1)))
    if not names:
        raise ValueError("La presentación no contiene diapositivas legibles.")
    blocks = []
    for name in names:
        for node in parts[name].iter():
            if node.tag == A + "tbl":
                blocks.append(("table", [["\n".join(_paragraph_text(p, A) for p in cell.iter(A + "p"))
                                           for cell in row.findall(A + "tc")] for row in node.findall(A + "tr")]))
            elif node.tag.endswith("}sp"):
                blocks.extend(("text", _paragraph_text(p, A)) for p in node.iter(A + "p"))
    return blocks


def _xlsx_blocks(content, warnings):
    blocks, count, formulas = [], 0, False
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in book:
            rows = []
            for row in sheet.iter_rows():
                count += len(row)
                if count > MAX_SHEET_CELLS:
                    warnings.append("La lectura de Excel se recortó a 50000 celdas; el original conserva todas las hojas.")
                    blocks.append(("table", rows))
                    return blocks
                values = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        formulas = True
                        value = None
                    if isinstance(value, (datetime, date)):
                        value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                    values.append(str(value).strip() if value is not None else "")
                if any(values):
                    rows.append(values)
            blocks.extend([("heading", sheet.title), ("table", rows)])
    finally:
        book.close()
        if formulas:
            warnings.append("Las fórmulas de Excel no se ejecutaron ni se propusieron como valores; revise esos campos en el original.")
    return blocks


def _read_blocks(name, content, warnings):
    suffix = PurePosixPath(name).suffix.lower()
    if suffix == ".json":
        warnings.append("El archivo JSON se conservó como soporte; sus datos no se aplican automáticamente a los formularios.")
        return [("raw", content.decode("utf-8-sig"))]
    if suffix in {".docx", ".xlsx", ".pptx"}:
        parts = _office_parts(content)
        if suffix == ".docx":
            return _docx_blocks(parts)
        if suffix == ".pptx":
            return _pptx_blocks(parts)
        return _xlsx_blocks(content, warnings)
    if suffix in {".txt", ".md", ".csv"}:
        text = content.decode("utf-8-sig")
        if suffix == ".csv":
            delimiter = ";" if text.partition("\n")[0].count(";") > text.partition("\n")[0].count(",") else ","
            return [("table", list(csv.reader(io.StringIO(text), delimiter=delimiter)))]
        return [("heading" if line.startswith("#") else "text", line.lstrip("# ")) for line in text.splitlines()]
    warnings.append("Se conservó únicamente el archivo original. La extracción de texto de este formato y el reconocimiento de texto en imágenes (OCR) no están habilitados.")
    return []


def _value(value, field, warnings):
    """Normaliza formatos explícitos; omite valores dudosos sin inferir hechos."""
    value = "\n".join(value) if isinstance(value, list) else str(value)
    value = value.strip()
    if not value or _normalize(value) in _EMPTY:
        return None
    kind = field["type"]
    if kind == "date" and re.fullmatch(r"\d{2}/\d{2}/\d{4}", value):
        day, month, year = value.split("/")
        value = f"{year}-{month}-{day}"
    elif kind == "number":
        if field.get("unit") == "COP":
            value = re.sub(r"(?i)\bCOP\b|\$|\s", "", value)
            if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", value):
                value = value.replace(".", "")
        if re.fullmatch(r"\d+,\d+", value):
            value = value.replace(",", ".")
    elif kind == "select":
        value = next((option["value"] for option in field["options"] if _normalize(option["label"]) == _normalize(value)), value)
    limit = 20000 if kind == "textarea" else 2000
    if len(value) > limit:
        warnings.append(f"El campo «{field['label']}» se recortó a {limit} caracteres; consulte el original completo.")
        value = value[:limit]
    try:
        result = normalize_value(value, field, field["label"])
        if field["key"] == "vigencia":
            if float(result) != int(float(result)):
                raise ValueError("La duración debe expresarse en meses enteros.")
            return int(float(result))
        if field["key"] == "objetivos_especificos":
            return [line.strip(" •-\t") for line in result.splitlines() if line.strip(" •-\t")]
        return result
    except (HTTPException, ValueError) as error:
        detail = error.detail if isinstance(error, HTTPException) else str(error)
        warnings.append(f"No se propuso «{field['label']}»: {detail}")
        return None


def _store(proposal, target, value, conflicts, warnings):
    group, field = target
    key = (group, field["key"])
    if value is None or key in conflicts:
        return
    previous = proposal[group].get(field["key"])
    if previous is not None and previous != value:
        proposal[group].pop(field["key"], None)
        conflicts.add(key)
        warnings.append(f"El archivo contiene valores diferentes para «{field['label']}»; el campo requiere revisión y no se propuso.")
    else:
        proposal[group][field["key"]] = value


def _table_values(rows, target, warnings):
    field = target[1]
    columns = {_normalize(label): column for column in field["columns"] for label in (column["label"], column["key"])}
    header = rows[0] if rows else []
    indexes = [(index, columns[_normalize(label)]) for index, label in enumerate(header) if _normalize(label) in columns]
    if len(indexes) < 2:
        return None
    values = []
    for row in rows[1:81]:
        data = {}
        for index, column in indexes:
            value = _value(row[index], column, warnings) if index < len(row) else None
            if value is not None:
                data[column["key"]] = value
        if data:
            try:
                values.append(validate_fields(data, field["columns"]))
            except HTTPException as error:
                warnings.append(f"No se propuso una fila de «{field['label']}»: {error.detail}")
    if len(rows) > 81:
        warnings.append(f"La tabla «{field['label']}» se recortó a 80 filas; el original conserva todas las filas.")
    return values or None


def _table_target(rows, definitions):
    """Asocia una tabla sin título solo cuando su encabezado identifica un único campo."""
    headers = {_normalize(value) for value in rows[0]} if rows else set()
    candidates = []
    for target in definitions:
        field = target[1]
        if field["type"] != "rows":
            continue
        column_names = {_normalize(value) for column in field["columns"] for value in (column["key"], column["label"])}
        matched = headers & column_names
        if len(matched) >= 2 and headers <= column_names:
            candidates.append(target)
    return candidates[0] if len(candidates) == 1 else None


def _propose(blocks, kind, warnings):
    proposal = {"comunes": {}, "borrador": {}, "proyecto": {}}
    definitions = [("comunes", field) for field in COMMON_FIELDS] + [("borrador", field) for field in DOCUMENT_DEFINITIONS.get(kind, {}).get("fields", [])] + [("proyecto", field) for field in PROJECT_FIELDS]
    labels = {_normalize(label): (group, field) for group, field in definitions for label in (field["label"], field["key"])}
    labels[_normalize("Título del Proyecto")] = ("proyecto", PROJECT_FIELDS[0])
    labels[_normalize("Objetivos específicos del proyecto")] = ("proyecto", PROJECT_FIELDS[3])
    conflicts, active, collected = set(), None, []
    # El bloque final permite vaciar la última sección sin un cierre especial.
    for block_kind, payload in blocks + [("heading", "")]:
        if block_kind == "raw":
            continue
        if block_kind == "table":
            table_target = active if active and active[1]["type"] == "rows" else _table_target(payload, definitions)
            if table_target:
                rows = _table_values(payload, table_target, warnings)
                _store(proposal, table_target, rows, conflicts, warnings)
                active = None
                continue
            if active and collected:
                _store(proposal, active, _value("\n".join(collected), active[1], warnings), conflicts, warnings)
            active, collected = None, []
            for row in payload:
                target = labels.get(_normalize(row[0])) if row else None
                if target and target[1]["type"] != "rows" and len(row) > 1:
                    _store(proposal, target, _value("\n".join(row[1:]), target[1], warnings), conflicts, warnings)
                elif len(row) == 1 and ":" in row[0]:
                    label, _, value = row[0].partition(":")
                    target = labels.get(_normalize(label))
                    if target and target[1]["type"] != "rows":
                        _store(proposal, target, _value(value, target[1], warnings), conflicts, warnings)
            continue
        text = payload.strip()
        label, separator, inline = text.partition(":")
        target = labels.get(_normalize(label if separator else text))
        if target or block_kind == "heading":
            if active and active[1]["type"] != "rows" and collected:
                _store(proposal, active, _value("\n".join(collected), active[1], warnings), conflicts, warnings)
            active, collected = target, []
            if target and separator and inline.strip() and target[1]["type"] != "rows":
                if target[1]["type"] == "textarea":
                    collected.append(inline.strip())
                else:
                    _store(proposal, target, _value(inline, target[1], warnings), conflicts, warnings)
                    active = None
        elif active and active[1]["type"] != "rows" and text:
            collected.append(text)
    for group, fields in (("comunes", COMMON_FIELDS), ("borrador", DOCUMENT_DEFINITIONS.get(kind, {}).get("fields", []))):
        try:
            proposal[group] = validate_fields(proposal[group], fields)
        except HTTPException as error:
            warnings.append(f"Los datos de «{group}» requieren revisión y no se propusieron: {error.detail}")
            proposal[group] = {}
    return proposal, conflicts


def _formulation_proposal(blocks, proposal, conflicts, warnings):
    """Reutiliza las secciones CAP conservando saltos de línea y avisando cada recorte."""
    rows = []
    for block_kind, payload in blocks:
        rows.extend(payload if block_kind == "table" else [[payload]])
    draft = {key: _section_content(rows, number) for number, key in GUIDED_SECTION_FIELDS.items()}
    general, specifics = _extract_objectives(rows)
    project = {"nombre": _find_labeled_value(rows, "Título del Proyecto") or _section_content(rows, 1),
               "objetivo_general": general, "objetivos_especificos": specifics}
    for group, data, definitions in (("proyecto", project, PROJECT_FIELDS), ("borrador", draft, DOCUMENT_DEFINITIONS["formulacion_proyecto"]["fields"])):
        fields = {field["key"]: field for field in definitions}
        for key, value in data.items():
            if value and key not in proposal[group] and (group, key) not in conflicts:
                _store(proposal, (group, fields[key]), _value(value, fields[key], warnings), conflicts, warnings)
    results = [line.strip(" -•\t") for line in _section_content(rows, 8).splitlines() if line.strip(" -•\t")]
    field = next(field for field in DOCUMENT_DEFINITIONS["formulacion_proyecto"]["fields"] if field["key"] == "resultados_esperados")
    if results and field["key"] not in proposal["borrador"] and ("borrador", field["key"]) not in conflicts:
        column = next(column for column in field["columns"] if column["key"] == "resultado")
        if len(results) > 80:
            warnings.append("La tabla «Resultados esperados» se recortó a 80 filas; el original conserva todas las filas.")
        values = [{"resultado": value} for item in results[:80] if (value := _value(item, column, warnings)) is not None]
        _store(proposal, ("borrador", field), values, conflicts, warnings)


def analyze_file(entry: dict, tipo: str | None = None) -> dict:
    """Devuelve texto y propuestas revisables sin guardar ni ejecutar archivos."""
    path, content = entry["ruta"], entry["content"]
    kind = _classify(path, tipo)
    warnings = []
    if kind == "documento_apoyo" and any(_normalize(part) == _normalize("6EvidenciasFotograficas") for part in PurePosixPath(path).parts[:-1]):
        warnings.append("El archivo no es una evidencia visual: se conserva como documento de apoyo y no completa la etapa de fotografías o videos.")
    period = None
    try:
        blocks = _read_blocks(path, content, warnings)
    except (ValueError, UnicodeError, KeyError, TypeError, OSError, OverflowError) as error:
        warnings.append(f"No se extrajo el texto: {error} Se conservó el archivo original.")
        blocks = []
    text = "\n".join("\n".join(" | ".join(row) for row in payload) if block_kind == "table" else payload for block_kind, payload in blocks).strip()
    if kind == "informe_bimensual":
        periods = {int(value) for value in re.findall(r"bimestr(?:e|al)[ _.:\-]*(\d{1,2})(?!\d)", path + "\n" + text, re.I)}
        if len(periods) == 1 and 1 <= next(iter(periods)) <= 30:
            period = next(iter(periods))
        else:
            warnings.append("No se identificó un bimestre único entre 1 y 30; indíquelo antes de completar el informe.")
    proposal, conflicts = _propose(blocks, kind, warnings)
    if blocks and kind == "formulacion_proyecto" and PurePosixPath(path).suffix.lower() == ".docx":
        _formulation_proposal(blocks, proposal, conflicts, warnings)
    if len(text) > MAX_EXTRACTED_TEXT:
        text = text[:MAX_EXTRACTED_TEXT]
        warnings.append(f"El texto extraído se recortó a {MAX_EXTRACTED_TEXT} caracteres; el archivo original se conserva completo.")
    if blocks and not any(proposal.values()):
        warnings.append("Se leyó el texto, pero no se identificaron campos con rótulos del formulario. Revise el contenido y registre los datos correspondientes.")
    return {"tipo": kind, "periodo_bimestre": period, "texto_extraido": text,
            "advertencias": list(dict.fromkeys(warnings)), "propuesta": proposal}
