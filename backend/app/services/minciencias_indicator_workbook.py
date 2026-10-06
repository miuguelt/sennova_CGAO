"""Construye una matriz descriptiva de datos existentes para seguimiento CTeI."""

import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CATEGORY_CODES = ("A", "B", "C", "D")
HEADER_FILL = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(bold=True, size=13, color="0F172A")
SUBTITLE_FONT = Font(size=9, color="64748B")


def _product_category(product):
    """Devuelve la categoría registrada o la que expresa un código A1-Dn."""
    category = str(getattr(product, "categoria", None) or "").strip().upper()
    if category in CATEGORY_CODES:
        return category

    product_type = str(getattr(product, "tipo", None) or "").strip().upper()
    match = re.match(r"^([ABCD]\d+)(?:\b|[\s\-:])", product_type)
    return match.group(1)[0] if match else "Sin categoría"


def _product_reporting_year(product):
    """Prioriza el año de reporte y usa el de publicación como alternativa."""
    reporting_year = getattr(product, "año_reporte", None)
    if isinstance(reporting_year, int) and not isinstance(reporting_year, bool):
        return reporting_year

    publication_date = getattr(product, "fecha_publicacion", None)
    return getattr(publication_date, "year", None)


def _product_support(product):
    """Indica si el registro tiene al menos un DOI o enlace de soporte."""
    return bool(str(getattr(product, "doi", None) or "").strip() or
                str(getattr(product, "url", None) or "").strip())


def _auto_fit_columns(worksheet, max_columns, max_width=48):
    """Ajusta el ancho de las columnas para facilitar lectura y filtros."""
    for column_index in range(1, max_columns + 1):
        column_letter = get_column_letter(column_index)
        longest = max(
            (len(str(cell.value)) for cell in worksheet[column_letter]
             if cell.value is not None),
            default=10,
        )
        worksheet.column_dimensions[column_letter].width = min(max(longest + 2, 12), max_width)


def _append_safe_row(worksheet, values):
    """Escribe texto externo como texto para evitar fórmulas al abrir el Excel."""
    safe_values = []
    for value in values:
        if isinstance(value, str) and value.lstrip(" \t\r\n")[:1] in {"=", "+", "-", "@"}:
            value = "'" + value
        safe_values.append(value)
    worksheet.append(safe_values)


def _style_table(worksheet, header_row, column_count):
    """Aplica encabezados y filtros a una tabla de la matriz."""
    for cell in worksheet[header_row][:column_count]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(column_count)}{worksheet.max_row}"
    _auto_fit_columns(worksheet, column_count)


def _profile_gaps(investigator):
    """Describe los campos de perfil sin diligenciar que afectan la lectura."""
    gaps = []
    if not str(getattr(investigator, "nivel_academico", None) or "").strip():
        gaps.append("Nivel académico")
    if not str(getattr(investigator, "cv_lac_url", None) or "").strip():
        gaps.append("Enlace CvLAC")
    if not str(getattr(investigator, "estado_cv_lac", None) or "").strip():
        gaps.append("Estado CvLAC")
    return ", ".join(gaps) if gaps else "Sin faltantes en estos campos"


def _build_researcher_rows(investigators, products_by_owner):
    """Calcula conteos descriptivos por investigador desde productos registrados."""
    rows = []
    for investigator in investigators:
        products = products_by_owner.get(str(investigator.id), [])
        category_counts = Counter(_product_category(product) for product in products)
        rows.append([
            investigator.nombre,
            investigator.nivel_academico or "",
            investigator.estado_cv_lac or "Sin dato",
            investigator.cv_lac_url or "",
            ", ".join(investigator.lineas_investigacion or []),
            len(products),
            category_counts["A"],
            category_counts["B"],
            category_counts["C"],
            category_counts["D"],
            category_counts["Sin categoría"],
            sum(1 for product in products if bool(product.is_verificado)),
            sum(1 for product in products if _product_support(product)),
            _profile_gaps(investigator),
        ])
    return rows


def build_minciencias_indicator_workbook(group, investigators, products, reporting_year=None):
    """Genera un Excel de seguimiento con procedencia, conteos y límites explícitos."""
    investigators = sorted(investigators, key=lambda item: (item.nombre or "").casefold())
    investigator_names = {str(item.id): item.nombre for item in investigators}
    products_by_owner = defaultdict(list)
    included_products = []
    for product in products:
        if str(product.owner_id) not in investigator_names:
            continue
        if reporting_year is not None and _product_reporting_year(product) != reporting_year:
            continue
        products_by_owner[str(product.owner_id)].append(product)
        included_products.append(product)

    included_products.sort(key=lambda product: (
        (investigator_names.get(str(product.owner_id)) or "").casefold(),
        (product.tipo or "").casefold(),
        (product.nombre or "").casefold(),
    ))

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary["A1"] = "MATRIZ DE SEGUIMIENTO DE INDICADORES MINCIENCIAS"
    summary.merge_cells("A1:B1")
    summary["A1"].font = TITLE_FONT
    summary["A1"].alignment = Alignment(horizontal="center")
    summary["A2"] = (
        f"Grupo: {group.nombre if group else 'No registrado'} | "
        f"Año: {reporting_year if reporting_year is not None else 'Todos'} | "
        f"Generado: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    summary.merge_cells("A2:B2")
    summary["A2"].font = SUBTITLE_FONT
    category_counts = Counter(_product_category(product) for product in included_products)
    summary_rows = [
        ("Investigadores incluidos", len(investigators)),
        ("Año del reporte", reporting_year if reporting_year is not None else "Todos"),
        ("Productos del período", len(included_products)),
        ("Productos verificados", sum(1 for item in included_products if bool(item.is_verificado))),
        ("Productos sin categoría", category_counts["Sin categoría"]),
        ("Productos sin soporte", sum(1 for item in included_products if not _product_support(item))),
        ("Clasificación registrada del grupo", group.clasificacion or "Sin dato" if group else "Sin grupo"),
    ]
    summary.append([])
    summary.append(["Indicador", "Valor"])
    for label, value in summary_rows:
        _append_safe_row(summary, [label, value])
    _style_table(summary, 4, 2)

    researcher_sheet = workbook.create_sheet("Investigadores")
    researcher_sheet.append(["MATRIZ POR INVESTIGADOR"])
    researcher_sheet.merge_cells("A1:N1")
    researcher_sheet["A1"].font = TITLE_FONT
    researcher_sheet.append([])
    researcher_sheet.append([
        "Investigador", "Nivel académico", "Estado CvLAC", "Enlace CvLAC",
        "Líneas de investigación", "Productos registrados", "Categoría A",
        "Categoría B", "Categoría C", "Categoría D", "Sin categoría",
        "Productos verificados", "Productos con DOI o URL", "Datos de perfil pendientes",
    ])
    for row in _build_researcher_rows(investigators, products_by_owner):
        _append_safe_row(researcher_sheet, row)
    _style_table(researcher_sheet, 3, 14)

    product_sheet = workbook.create_sheet("Productos")
    product_sheet.append(["DETALLE DE PRODUCTOS REGISTRADOS"])
    product_sheet.merge_cells("A1:J1")
    product_sheet["A1"].font = TITLE_FONT
    product_sheet.append([])
    product_sheet.append([
        "Investigador", "Producto", "Código/tipo registrado", "Categoría registrada o derivada",
        "Año de reporte usado", "Verificado en la aplicación", "Proyecto asociado",
        "DOI", "URL o soporte", "Tiene DOI o URL",
    ])
    for product in included_products:
        project = getattr(product, "proyecto", None)
        project_name = (project.nombre_corto or project.nombre) if project else "Sin proyecto asociado"
        _append_safe_row(product_sheet, [
            investigator_names.get(str(product.owner_id), ""),
            product.nombre,
            product.tipo,
            _product_category(product),
            _product_reporting_year(product) or "",
            "Sí" if product.is_verificado else "No",
            project_name,
            product.doi or "",
            product.url or "",
            "Sí" if _product_support(product) else "No",
        ])
    _style_table(product_sheet, 3, 10)

    methodology = workbook.create_sheet("Metodología")
    methodology.append(["Criterio", "Descripción"])
    methodology_rows = [
        ("Procedencia", "Datos existentes en los perfiles y productos asociados a integrantes activos del grupo."),
        ("Año", "Se usa el año de reporte registrado; si no existe, se usa el año de publicación."),
        ("Categoría", "Se usa la categoría A-D registrada; si falta, se deriva únicamente de códigos tipo A1, B1, C1 o D1."),
        ("Verificación", "Refleja el estado de verificación guardado en esta aplicación, no una validación de MinCiencias."),
        ("Soporte", "Se marca disponible cuando el producto tiene al menos un DOI o una URL."),
        ("Alcance", "No calcula ni certifica la clasificación de un investigador o grupo ni otorga puntajes."),
        ("Uso", "Revise los datos y la versión vigente del modelo y la convocatoria de MinCiencias antes de tomar decisiones."),
    ]
    for row in methodology_rows:
        methodology.append(row)
    _style_table(methodology, 1, 2)
    methodology.column_dimensions["B"].width = 100
    for row in methodology.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
