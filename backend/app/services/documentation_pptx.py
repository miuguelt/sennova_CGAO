"""Adaptador de presentaciones y pósteres editables con paginación del contenido."""

import io

from docx.shared import Cm
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.util import Inches, Pt as PptxPt

from app.services.documentation_catalog import COMMON_FIELDS, DOCUMENT_DEFINITIONS
from app.services.documentation_rendering_values import PENDING, field_text, format_value, text_chunks

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _slide_text(slide, text, left, top, width, height, size, *, bold=False, font="Arial"):
    """Agrega texto editable en un área definida con tamaño legible."""
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.08)
    frame.margin_top = frame.margin_bottom = Inches(0.04)
    for index, line in enumerate(str(text).splitlines()):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = PptxPt(size)
        paragraph.font.name = font
        paragraph.font.bold = bold
        paragraph.font.color.rgb = PptxRGBColor(0, 0, 0)
        paragraph.space_after = PptxPt(5)
    return shape



def _deck_slide(presentation, title, text, context):
    """Crea páginas de presentación cuando el contenido supera una diapositiva."""
    for index, chunk in enumerate(text_chunks(text)):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        label = title if index == 0 else f"{title} continuación {index + 1}"
        _slide_text(slide, label, Inches(0.65), Inches(0.45), Inches(12), Inches(1.2), 30, bold=True)
        _slide_text(slide, chunk, Inches(0.65), Inches(1.75), Inches(12), Inches(4.85), 21)
        _slide_text(slide, f"Borrador para revisión · Versión {context.get('version', 1)}", Inches(0.65), Inches(6.9), Inches(12), Inches(0.3), 10)



def _deck_table(presentation, field, values, context):
    """Distribuye tablas editables entre diapositivas y conserva celdas extensas."""
    columns = field["columns"]
    records = []
    for row in values.get(field["key"]) or [{}]:
        chunks = [text_chunks(format_value(row.get(column["key"]), column), width=max(18, 65 // len(columns)), lines=4) for column in columns]
        for index in range(max(len(parts) for parts in chunks)):
            records.append([parts[index] if index < len(parts) else "" for parts in chunks])
    for start in range(0, len(records), 3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = field["label"] + (f" continuación {start // 3 + 1}" if start else "")
        _slide_text(slide, title, Inches(0.65), Inches(0.45), Inches(12), Inches(1.1), 30, bold=True)
        rows = records[start:start + 3]
        table = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(0.65), Inches(1.8), Inches(12), Inches(4.65)).table
        for column_index, column in enumerate(columns):
            table.cell(0, column_index).text = column["label"]
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.margin_left = cell.margin_right = Inches(0.1)
                cell.margin_top = cell.margin_bottom = Inches(0.06)
                if row_index == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = PptxRGBColor(232, 238, 243)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = PptxPt(16)
                    paragraph.font.name = "Arial"
                    paragraph.font.bold = row_index == 0
                    paragraph.font.color.rgb = PptxRGBColor(0, 0, 0)
        _slide_text(slide, f"Borrador para revisión · Versión {context.get('version', 1)}", Inches(0.65), Inches(6.9), Inches(12), Inches(0.3), 10)



def _deck_schedule(presentation, field, common, context):
    """Separa la tabla base del cronograma de sus datos adicionales para mantenerlos legibles."""
    columns = field["columns"]
    activity_column = next(column for column in columns if column["key"] == "actividad")
    base_field = dict(field, columns=[column for column in columns if not column.get("optional_detail")])
    _deck_table(presentation, base_field, common, context)

    detail_columns = [column for column in columns if column.get("optional_detail")]
    detail_key = "programacion_gantt"
    detail_rows = []
    for row in common.get(field["key"]) or []:
        details = [f"{column['label']}: {format_value(row.get(column['key']), column)}"
                   for column in detail_columns if row.get(column["key"]) not in (None, "")]
        if details:
            detail_rows.append({
                "actividad": format_value(row.get("actividad"), activity_column),
                detail_key: "\n".join(details),
            })
    if detail_rows:
        detail_field = {
            "key": "programacion_gantt", "label": "Programación para el diagrama de Gantt",
            "columns": [
                activity_column,
                {"key": detail_key, "label": "Fase, fechas, horario y lugar", "type": "textarea"},
            ],
        }
        _deck_table(presentation, detail_field, {detail_key: detail_rows}, context)



def render_poster(context, common, data, photos):
    """Mantiene el póster vertical y crea láminas adicionales si el contenido crece."""
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Cm(90), Cm(120)
    sections = [(field["label"], field_text(field, data)) for field in DOCUMENT_DEFINITIONS["poster_producto"]["fields"]]
    sections.insert(3, ("Objetivos", str(context.get("objetivo_general") or PENDING) + "\n" + "\n".join(context.get("objetivos_especificos") or [])))
    authors = "; ".join(str(row.get("nombre") or PENDING) for row in common.get("equipo", [])) or PENDING
    attribution = "\n".join([
        str(common.get("centro") or PENDING), str(common.get("regional") or PENDING),
        f"Responsable: {common.get('responsable') or PENDING}", f"Autoría: {authors}",
        f"Código SGPS: {context.get('codigo_sgps') or PENDING}",
        f"Producto: {(context.get('producto') or {}).get('nombre') or PENDING}",
    ])
    attribution_parts = text_chunks(attribution, width=100, lines=6)
    sections.extend(("Autoría y filiación (continuación)", part) for part in attribution_parts[1:])
    if common.get("inconsistencias_fuente"):
        sections.append(("Datos pendientes de validación", common["inconsistencias_fuente"]))
    records = [(title, chunk) for title, text in sections for chunk in text_chunks(text, width=55, lines=8)]
    remaining_photos = list(photos)
    for start in range(0, len(records), 10):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _slide_text(slide, context["nombre"], Cm(4), Cm(3), Cm(82), Cm(12), 68, bold=True, font="Work Sans")
        _slide_text(slide, attribution_parts[0], Cm(4), Cm(16), Cm(82), Cm(7), 26, font="Work Sans")
        for index, (title, text) in enumerate(records[start:start + 10]):
            left = Cm(4 + (index % 2) * 43)
            top = Cm(25 + (index // 2) * 17)
            _slide_text(slide, title, left, top, Cm(39), Cm(3), 40, bold=True, font="Work Sans")
            _slide_text(slide, text, left, top + Cm(3.5), Cm(39), Cm(12), 34, font="Work Sans")
        if start == 0 and len(records) < 10 and remaining_photos and len(text_chunks(remaining_photos[0][1], width=60, lines=3)) == 1:
            content, caption = remaining_photos.pop(0)
            slot = len(records)
            left, top = Cm(4 + (slot % 2) * 43), Cm(25 + (slot // 2) * 17)
            _slide_text(slide, "Registro fotográfico", left, top, Cm(39), Cm(3), 40, bold=True, font="Work Sans")
            with Image.open(io.BytesIO(content)) as image:
                ratio = image.width / image.height
            width = min(Cm(37), int(Cm(9) * ratio))
            slide.shapes.add_picture(io.BytesIO(content), left, top + Cm(3.5), width=width, height=int(width / ratio))
            _slide_text(slide, caption, left, top + Cm(12.8), Cm(39), Cm(4), 30, font="Work Sans")
        _slide_text(slide, f"Borrador para revisión · Versión {context.get('version', 1)}", Cm(4), Cm(115), Cm(82), Cm(2), 24, font="Work Sans")
    for content, caption in remaining_photos:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _slide_text(slide, "Registro fotográfico del producto", Cm(4), Cm(4), Cm(82), Cm(6), 64, bold=True)
        with Image.open(io.BytesIO(content)) as image:
            ratio = image.width / image.height
        width = min(Cm(80), int(Cm(90) * ratio))
        height = int(width / ratio)
        slide.shapes.add_picture(io.BytesIO(content), int((presentation.slide_width - width) / 2), Cm(15), width=width, height=height)
        _slide_text(slide, caption, Cm(4), Cm(108), Cm(82), Cm(6), 32)
        _slide_text(slide, f"Borrador para revisión · Versión {context.get('version', 1)}", Cm(4), Cm(115), Cm(82), Cm(2), 24)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue(), PPTX_MIME



def render_presentation(context, common, data):
    """Conserva la secuencia temática de la presentación de referencia."""
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.3333), Inches(7.5)
    _deck_slide(presentation, "Presentación del proyecto", f"{context['nombre']}\nCódigo SGPS {context.get('codigo_sgps') or PENDING}\n{common.get('centro') or PENDING}\n{common.get('responsable') or PENDING}", context)
    required_common = DOCUMENT_DEFINITIONS["presentacion_proyecto"]["required_common"]
    context_lines = [f"{field['label']}: {format_value(common.get(field['key']), field)}" for field in COMMON_FIELDS
                     if field["type"] != "rows" and field["key"] != "inconsistencias_fuente"
                     and (common.get(field["key"]) or field["key"] in required_common)]
    context_lines.append("Presupuesto total: " + format_value(context.get("presupuesto_total"), {"label": "Presupuesto total", "type": "number", "unit": "COP"}))
    _deck_slide(presentation, "Información general del proyecto", "\n".join(context_lines), context)
    for key in ("equipo", "presupuesto"):
        field = next(item for item in COMMON_FIELDS if item["key"] == key)
        public_field = dict(field, columns=[column for column in field["columns"] if not column.get("optional_detail")])
        _deck_table(presentation, public_field, common, context)
    for field in DOCUMENT_DEFINITIONS["presentacion_proyecto"]["fields"]:
        if field["type"] == "rows":
            _deck_table(presentation, field, data, context)
        else:
            _deck_slide(presentation, field["label"], field_text(field, data), context)
        if field["key"] == "justificacion":
            _deck_slide(presentation, "Objetivo general", context.get("objetivo_general") or PENDING, context)
            objectives = context.get("objetivos_especificos") or []
            _deck_slide(presentation, "Objetivos específicos", "\n".join(objectives) if isinstance(objectives, list) else objectives, context)
        if field["key"] == "fases":
            schedule = next(item for item in COMMON_FIELDS if item["key"] == "cronograma")
            _deck_schedule(presentation, schedule, common, context)
    if common.get("inconsistencias_fuente"):
        _deck_slide(presentation, "Datos pendientes de validación", common["inconsistencias_fuente"], context)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue(), PPTX_MIME

